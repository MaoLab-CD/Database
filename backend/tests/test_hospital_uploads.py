from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from fastapi import HTTPException

from app.api.routes import hospital_uploads as hospital_upload_routes
from app.api.routes.hospital_uploads import resolve_stored_path, verify_stored_file
from app.core.config import settings
from app.services.hospital_file_storage import (
    ENCRYPTED_FILE_MAGIC,
    decrypt_hospital_file,
    encrypt_hospital_file,
    sha256_bytes,
    temporary_plaintext_file,
    write_encrypted_hospital_file,
)
from app.services.hospital_uploads import (
    build_hospital_review_preview,
    build_hospital_review_preview_workbook,
    build_hospital_submission_preview,
    build_hospital_upload_template,
    inspect_xlsx_package,
    resolve_workbook_sheet_name,
    sha256_file,
    validate_hospital_upload,
)
from app.services.sample_code_excel import SampleCodeAssignment


@pytest.fixture
def workbook_path() -> Path:
    path = Path(__file__).parent / f"_hospital_upload_{uuid4().hex}.xlsx"
    try:
        yield path
    finally:
        path.unlink(missing_ok=True)


class _Result:
    def __init__(self, value=None, values=None):
        self.value = value
        self.values = values or []

    def scalar_one_or_none(self):
        return self.value

    def scalars(self):
        return iter(self.values)


class _ValidationSession:
    def execute(self, statement, params=None):
        sql = str(statement)
        if "FROM centers" in sql:
            return _Result(value=1)
        if "FROM specimen_types" in sql:
            return _Result(values=["全血", "血清", "血浆", "组织", "粪便", "DNA", "RNA", "干粉"])
        if "FROM samples" in sql:
            return _Result(values=[])
        raise AssertionError(f"unexpected SQL: {sql}")


def _write_valid_upload(path: Path) -> Path:
    content = build_hospital_upload_template()
    workbook = load_workbook(BytesIO(content))
    worksheet = workbook["样本数据"]
    worksheet["A2"] = "HOSP-0001"
    worksheet["B2"] = "2026-08-07 09:30:00"
    worksheet["C2"] = "全血"
    workbook.save(path)
    workbook.close()
    return path


def test_template_is_valid_xlsx_and_contains_required_columns(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)
    inspect_xlsx_package(path)

    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook.sheetnames == ["样本数据", "填写说明"]
        headers = [cell.value for cell in next(workbook["样本数据"].iter_rows())]
        assert "序" in headers
        assert "采集时间" in headers
        assert "标本" in headers
        assert "sample_code" not in headers
    finally:
        workbook.close()


def test_valid_upload_accepts_hospital_original_fields_without_sample_code(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)
    result = validate_hospital_upload(
        _ValidationSession(),
        path,
        "样本数据",
        "510115-001",
    )
    assert result.is_valid
    assert result.total_rows == 1
    assert result.new_rows == 1
    assert result.summary["specimen_type_counts"] == {"全血": 1}


def test_review_preview_keeps_original_fields_and_appends_candidate_code(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)
    assignment = SampleCodeAssignment(
        row_number=2,
        sample_id="HOSP-0001",
        collection_date="2026-08-07",
        specimen_type="全血",
        sample_code="510115-001-2608070001",
    )

    headers, rows, sensitive_columns, total_rows = build_hospital_review_preview(
        path,
        "样本数据",
        [assignment],
    )

    assert total_rows == 1
    assert headers[headers.index("序") + 1] == "sample_code"
    assert headers[-2:] == ["校验结果", "校验说明"]
    assert rows[0]["row_number"] == 2
    assert rows[0]["values"][headers.index("序")] == "HOSP-0001"
    assert rows[0]["values"][headers.index("sample_code")] == assignment.sample_code
    assert rows[0]["values"][headers.index("校验结果")] == "通过"
    assert headers.index("姓名") in sensitive_columns


def test_submission_preview_keeps_original_fields_without_platform_code(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)

    headers, rows, sensitive_columns, total_rows = build_hospital_submission_preview(
        path,
        "样本数据",
    )

    assert total_rows == 1
    assert "sample_code" not in headers
    assert rows[0]["row_number"] == 2
    assert rows[0]["values"][headers.index("序")] == "HOSP-0001"
    assert headers.index("姓名") in sensitive_columns


def test_review_preview_workbook_contains_full_rows_and_review_notes(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)
    assignment = SampleCodeAssignment(
        row_number=2,
        sample_id="HOSP-0001",
        collection_date="2026-08-07",
        specimen_type="全血",
        sample_code="510115-001-2608070001",
    )

    content = build_hospital_review_preview_workbook(path, "样本数据", [assignment])
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["审核预览", "审核说明"]
        rows = list(workbook["审核预览"].iter_rows(values_only=True))
        headers = list(rows[0])
        assert rows[1][headers.index("序")] == "HOSP-0001"
        assert headers[headers.index("序") + 1] == "sample_code"
        assert rows[1][headers.index("sample_code")] == assignment.sample_code
        assert rows[1][headers.index("校验结果")] == "通过"
    finally:
        workbook.close()


def test_review_preview_keeps_generated_sample_code_unique(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)
    workbook = load_workbook(path)
    worksheet = workbook["样本数据"]
    worksheet["D1"] = "sample_code"
    worksheet["D2"] = "510115-001-legacy"
    workbook.save(path)
    workbook.close()
    assignment = SampleCodeAssignment(
        row_number=2,
        sample_id="HOSP-0001",
        collection_date="2026-08-07",
        specimen_type="全血",
        sample_code="510115-001-2608070001",
    )

    headers, rows, _, _ = build_hospital_review_preview(path, "样本数据", [assignment])

    assert headers.count("sample_code") == 1
    assert "原表_sample_code" in headers
    assert rows[0]["values"][headers.index("sample_code")] == assignment.sample_code
    assert rows[0]["values"][headers.index("原表_sample_code")] == "510115-001-legacy"


def test_single_sheet_is_auto_selected_when_requested_name_is_different(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)
    workbook = load_workbook(path)
    workbook.remove(workbook["填写说明"])
    workbook["样本数据"].title = "Sheet2"
    workbook.save(path)
    workbook.close()

    assert resolve_workbook_sheet_name(path, "样本数据") == "Sheet2"


def test_existing_platform_code_must_match_account_center(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)
    workbook = load_workbook(path)
    worksheet = workbook["样本数据"]
    worksheet["D1"] = "sample_code"
    worksheet["D2"] = "510115-001-2608070001"
    workbook.save(path)
    workbook.close()

    result = validate_hospital_upload(
        _ValidationSession(), path, "样本数据", "510105-001"
    )
    assert not result.is_valid
    assert result.error_rows == 1
    assert "与当前账号绑定中心 510105-001 不一致" in result.errors[0]["reason"]


def test_formula_cells_are_rejected(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)
    workbook = load_workbook(path)
    workbook["样本数据"]["D2"] = "=1+1"
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="公式"):
        validate_hospital_upload(_ValidationSession(), path, "样本数据", "510115-001")


def test_csv_upload_is_supported_and_uses_csv_sheet(tmp_path: Path) -> None:
    path = tmp_path / "hospital_upload.csv"
    path.write_text(
        "序,采集时间,标本,额外字段\nHOSP-CSV-1,2026-08-08 10:00:00,全血,保留内容\n",
        encoding="utf-8-sig",
    )

    assert resolve_workbook_sheet_name(path, "任意工作表") == "CSV"
    result = validate_hospital_upload(
        _ValidationSession(), path, "CSV", "510115-001"
    )
    assert result.is_valid
    assert result.total_rows == 1


def test_csv_formula_content_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "hospital_formula.csv"
    path.write_text(
        "序,采集时间,标本,备注\nHOSP-CSV-2,2026-08-08,全血,=1+1\n",
        encoding="utf-8-sig",
    )

    with pytest.raises(ValueError, match="公式内容"):
        validate_hospital_upload(_ValidationSession(), path, "CSV", "510115-001")


def test_macro_free_xlsm_container_is_supported(workbook_path: Path) -> None:
    xlsm_path = workbook_path.with_suffix(".xlsm")
    try:
        _write_valid_upload(workbook_path)
        xlsm_path.write_bytes(workbook_path.read_bytes())
        result = validate_hospital_upload(
            _ValidationSession(), xlsm_path, "样本数据", "510115-001"
        )
        assert result.is_valid
    finally:
        xlsm_path.unlink(missing_ok=True)


def test_embedded_objects_are_rejected(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)
    with ZipFile(path, "a", ZIP_DEFLATED) as archive:
        archive.writestr("xl/embeddings/object1.bin", b"unsafe")

    with pytest.raises(ValueError, match="嵌入对象"):
        inspect_xlsx_package(path)


def test_stored_file_hash_must_match_before_submit_or_review(workbook_path) -> None:
    path = _write_valid_upload(workbook_path)
    expected_hash = sha256_file(path)
    verify_stored_file(path, expected_hash)

    path.write_bytes(path.read_bytes() + b"changed")
    with pytest.raises(HTTPException) as exc_info:
        verify_stored_file(path, expected_hash)
    assert exc_info.value.status_code == 409


def test_hospital_upload_encryption_round_trip_hides_plaintext() -> None:
    plaintext = b"patient-id=510115-001;phone=13800000000"

    encrypted = encrypt_hospital_file(plaintext)
    decrypted, was_encrypted = decrypt_hospital_file(encrypted)

    assert encrypted.startswith(ENCRYPTED_FILE_MAGIC)
    assert plaintext not in encrypted
    assert b"13800000000" not in encrypted
    assert was_encrypted is True
    assert decrypted == plaintext


def test_dedicated_key_can_be_enabled_after_secret_key_fallback(monkeypatch) -> None:
    monkeypatch.setattr(settings, "hospital_upload_encryption_key", None)
    encrypted_with_fallback = encrypt_hospital_file(b"legacy encrypted upload")

    monkeypatch.setattr(settings, "hospital_upload_encryption_key", "new-dedicated-test-key")
    decrypted, was_encrypted = decrypt_hospital_file(encrypted_with_fallback)

    assert was_encrypted is True
    assert decrypted == b"legacy encrypted upload"


def test_encrypted_stored_file_hash_is_checked_on_plaintext(tmp_path: Path) -> None:
    plaintext = b"sensitive hospital upload"
    path = tmp_path / "hospital-upload.xlsx.enc"
    write_encrypted_hospital_file(path, plaintext)

    verify_stored_file(path, sha256_bytes(plaintext))

    encrypted = bytearray(path.read_bytes())
    encrypted[-1] ^= 1
    path.write_bytes(encrypted)
    with pytest.raises(HTTPException) as exc_info:
        verify_stored_file(path, sha256_bytes(plaintext))
    assert exc_info.value.status_code == 409


def test_temporary_plaintext_file_is_removed_after_use() -> None:
    with temporary_plaintext_file(b"private", ".xlsx") as path:
        directory = path.parent
        assert path.read_bytes() == b"private"
        assert path.is_file()

    assert not path.exists()
    assert not directory.exists()


def test_stored_file_path_survives_deployment_directory_change(tmp_path: Path, monkeypatch) -> None:
    quarantine = tmp_path / "quarantine"
    quarantine.mkdir()
    stored = quarantine / "random-upload.xlsx"
    stored.write_bytes(b"test")
    monkeypatch.setattr(hospital_upload_routes, "HOSPITAL_UPLOAD_DIR", quarantine)

    legacy_absolute_path = Path("C:/old-deployment/uploads/hospital/quarantine/random-upload.xlsx")
    assert resolve_stored_path(str(legacy_absolute_path)) == stored.resolve()
