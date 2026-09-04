from __future__ import annotations

import os
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Connection, text
from sqlalchemy.orm import Session

from app.api.deps import CurrentUser, get_current_user
from app.api.routes.hospital_uploads import HOSPITAL_UPLOAD_DIR
from app.db.session import engine, get_db
from app.main import app
from app.services.hospital_uploads import build_hospital_upload_template


RUN_DB_INTEGRATION_TESTS = os.getenv("RUN_DB_INTEGRATION_TESTS") == "1"
pytestmark = pytest.mark.skipif(
    not RUN_DB_INTEGRATION_TESTS,
    reason="设置 RUN_DB_INTEGRATION_TESTS=1 后才连接 PostgreSQL 测试库",
)


@dataclass
class UploadIntegrationContext:
    connection: Connection
    outer_transaction: object
    session: Session
    client: TestClient
    hospital_user: CurrentUser
    admin_user: CurrentUser
    active_user: dict[str, CurrentUser]
    center_code: str
    stored_files: list[Path]


@pytest.fixture
def upload_context() -> UploadIntegrationContext:
    connection = engine.connect()
    outer_transaction = connection.begin()
    session = Session(
        bind=connection,
        autoflush=False,
        expire_on_commit=False,
        join_transaction_mode="create_savepoint",
    )
    suffix = uuid4().hex[:12]
    center_code = session.execute(
        text("SELECT center_code FROM centers WHERE is_active = true ORDER BY center_code LIMIT 1")
    ).scalar_one()
    hospital_id = session.execute(
        text(
            """
            INSERT INTO users (
                username, password_hash, display_name, role, is_super_admin,
                permissions, status, account_type
            ) VALUES (
                :username, 'integration-test-only', '医院上传集成测试', 'user', false,
                '{"hospital_data_submit": true}'::jsonb, 'active', 'hospital'
            ) RETURNING id
            """
        ),
        {"username": f"it-hospital-{suffix}"},
    ).scalar_one()
    admin_id = session.execute(
        text(
            """
            INSERT INTO users (
                username, password_hash, display_name, role, is_super_admin,
                permissions, status, account_type
            ) VALUES (
                :username, 'integration-test-only', '上传审核集成测试', 'admin', false,
                '{}'::jsonb, 'active', 'internal'
            ) RETURNING id
            """
        ),
        {"username": f"it-reviewer-{suffix}"},
    ).scalar_one()
    session.execute(
        text("INSERT INTO user_center_scopes (user_id, center_code) VALUES (:user_id, :center_code)"),
        {"user_id": hospital_id, "center_code": center_code},
    )
    session.commit()

    hospital_user = CurrentUser(
        id=hospital_id,
        username=f"it-hospital-{suffix}",
        display_name="医院上传集成测试",
        role="user",
        is_super_admin=False,
        permissions={"hospital_data_submit": True},
        password_reset_required=False,
        account_type="hospital",
        center_codes=(center_code,),
    )
    admin_user = CurrentUser(
        id=admin_id,
        username=f"it-reviewer-{suffix}",
        display_name="上传审核集成测试",
        role="admin",
        is_super_admin=False,
        permissions={},
        password_reset_required=False,
    )
    active_user = {"value": hospital_user}
    stored_files: list[Path] = []

    def override_db():
        yield session

    def override_user() -> CurrentUser:
        return active_user["value"]

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = override_user
    client = TestClient(app)

    try:
        yield UploadIntegrationContext(
            connection=connection,
            outer_transaction=outer_transaction,
            session=session,
            client=client,
            hospital_user=hospital_user,
            admin_user=admin_user,
            active_user=active_user,
            center_code=center_code,
            stored_files=stored_files,
        )
    finally:
        client.close()
        app.dependency_overrides.clear()
        for path in stored_files:
            path.unlink(missing_ok=True)
        session.close()
        if outer_transaction.is_active:
            outer_transaction.rollback()
        connection.close()


def _build_upload() -> bytes:
    workbook = load_workbook(BytesIO(build_hospital_upload_template()))
    worksheet = workbook["样本数据"]
    worksheet["A2"] = "900001"
    worksheet["B2"] = "2099-01-01 09:30:00"
    worksheet["C2"] = "全血"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_hospital_upload_requires_review_before_atomic_import(
    upload_context: UploadIntegrationContext,
) -> None:
    max_serial = upload_context.session.execute(
        text(
            """
            SELECT COALESCE(max(substring(sample_seq FROM 7 FOR 4)::int), 0)
            FROM samples
            WHERE is_deleted = false
              AND center_code = :center_code
              AND sample_seq ~ '^990101[0-9]{4}d?$'
            """
        ),
        {"center_code": upload_context.center_code},
    ).scalar_one()
    sample_code = f"{upload_context.center_code}-990101{int(max_serial) + 1:04d}"
    upload_response = upload_context.client.post(
        "/api/hospital-uploads/upload",
        data={"sheet_name": "样本数据"},
        files={
            "file": (
                "hospital-data.xlsx",
                _build_upload(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert upload_response.status_code == 200, upload_response.text
    batch = upload_response.json()["batch"]
    assert batch["status"] == "validated"
    assert [event["event_type"] for event in upload_response.json()["events"]] == ["uploaded"]
    stored_name = upload_context.session.execute(
        text("SELECT file_path FROM hospital_upload_batches WHERE id = :id"),
        {"id": batch["id"]},
    ).scalar_one()
    stored_path = HOSPITAL_UPLOAD_DIR / Path(stored_name).name
    upload_context.stored_files.append(stored_path)
    stored_content = stored_path.read_bytes()
    assert stored_name.endswith(".enc")
    assert not stored_content.startswith(b"PK")
    assert b"900001" not in stored_content
    assert upload_context.session.execute(
        text("SELECT count(*) FROM samples WHERE sample_code = :sample_code"),
        {"sample_code": sample_code},
    ).scalar_one() == 0

    submission_preview_response = upload_context.client.get(
        f"/api/hospital-uploads/{batch['id']}/submission-preview"
    )
    assert submission_preview_response.status_code == 200, submission_preview_response.text
    submission_preview = submission_preview_response.json()
    assert submission_preview["total_rows"] == 1
    assert "序" in submission_preview["headers"]
    assert "sample_code" not in submission_preview["headers"]
    assert submission_preview["rows"][0]["values"][submission_preview["headers"].index("序")] == "900001"

    submit_response = upload_context.client.post(
        f"/api/hospital-uploads/{batch['id']}/submit"
    )
    assert submit_response.status_code == 200
    assert submit_response.json()["batch"]["status"] == "submitted"

    upload_context.active_user["value"] = upload_context.admin_user
    preview_response = upload_context.client.get(
        f"/api/hospital-uploads/{batch['id']}/code-preview"
    )
    assert preview_response.status_code == 200, preview_response.text
    assert preview_response.json()["rows"][0] == {
        "row_number": 2,
        "sample_id": "900001",
        "collection_date": "2099-01-01",
        "specimen_type": "全血",
        "sample_code": sample_code,
    }
    review_preview_response = upload_context.client.get(
        f"/api/hospital-uploads/{batch['id']}/review-preview"
    )
    assert review_preview_response.status_code == 200, review_preview_response.text
    review_preview = review_preview_response.json()
    assert review_preview["total_rows"] == 1
    assert "序" in review_preview["headers"]
    assert review_preview["headers"][review_preview["headers"].index("序") + 1] == "sample_code"
    candidate_index = review_preview["headers"].index("sample_code")
    assert review_preview["rows"][0]["values"][candidate_index] == sample_code

    review_workbook_response = upload_context.client.get(
        f"/api/hospital-uploads/{batch['id']}/review-preview.xlsx"
    )
    assert review_workbook_response.status_code == 200, review_workbook_response.text
    assert review_workbook_response.headers["cache-control"] == "no-store"
    review_workbook = load_workbook(BytesIO(review_workbook_response.content), read_only=True, data_only=True)
    try:
        exported_rows = list(review_workbook["审核预览"].iter_rows(values_only=True))
        exported_headers = list(exported_rows[0])
        assert exported_rows[1][exported_headers.index("序")] == "900001"
        assert exported_headers[exported_headers.index("序") + 1] == "sample_code"
        assert exported_rows[1][exported_headers.index("sample_code")] == sample_code
    finally:
        review_workbook.close()

    approve_response = upload_context.client.post(
        f"/api/hospital-uploads/{batch['id']}/review",
        json={"action": "approve", "comment": "校验通过"},
    )
    assert approve_response.status_code == 200, approve_response.text
    assert approve_response.json()["batch"]["status"] == "approved"
    assert [event["event_type"] for event in approve_response.json()["events"]] == [
        "uploaded",
        "submission_preview_viewed",
        "submitted",
        "review_preview_viewed",
        "review_preview_downloaded",
        "approved",
    ]
    assert upload_context.session.execute(
        text("SELECT count(*) FROM samples WHERE sample_code = :sample_code"),
        {"sample_code": sample_code},
    ).scalar_one() == 0

    import_response = upload_context.client.post(
        f"/api/hospital-uploads/{batch['id']}/import",
        json={"default_status": "not_stored"},
    )
    assert import_response.status_code == 200, import_response.text
    assert import_response.json()["batch"]["status"] == "imported"
    assert [event["event_type"] for event in import_response.json()["events"]] == [
        "uploaded",
        "submission_preview_viewed",
        "submitted",
        "review_preview_viewed",
        "review_preview_downloaded",
        "approved",
        "imported",
    ]
    stored = upload_context.session.execute(
        text("SELECT center_code, sample_status FROM samples WHERE sample_code = :sample_code"),
        {"sample_code": sample_code},
    ).mappings().one()
    assert stored == {"center_code": upload_context.center_code, "sample_status": "not_stored"}


def test_rejected_upload_requires_a_new_file(
    upload_context: UploadIntegrationContext,
) -> None:
    upload_response = upload_context.client.post(
        "/api/hospital-uploads/upload",
        data={"sheet_name": "样本数据"},
        files={
            "file": (
                "hospital-data.xlsx",
                _build_upload(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert upload_response.status_code == 200, upload_response.text
    batch = upload_response.json()["batch"]
    stored_name = upload_context.session.execute(
        text("SELECT file_path FROM hospital_upload_batches WHERE id = :id"),
        {"id": batch["id"]},
    ).scalar_one()
    upload_context.stored_files.append(HOSPITAL_UPLOAD_DIR / Path(stored_name).name)

    submit_response = upload_context.client.post(
        f"/api/hospital-uploads/{batch['id']}/submit"
    )
    assert submit_response.status_code == 200, submit_response.text

    upload_context.active_user["value"] = upload_context.admin_user
    reject_response = upload_context.client.post(
        f"/api/hospital-uploads/{batch['id']}/review",
        json={"action": "reject", "comment": "请修改后重新上传"},
    )
    assert reject_response.status_code == 200, reject_response.text
    assert reject_response.json()["batch"]["status"] == "rejected"

    upload_context.active_user["value"] = upload_context.hospital_user
    resubmit_response = upload_context.client.post(
        f"/api/hospital-uploads/{batch['id']}/submit"
    )
    assert resubmit_response.status_code == 409
    assert "重新上传" in resubmit_response.json()["detail"]
