'''
Author: 袁瑞 && 2502099390@qq.com
Date: 2026-06-12 13:39:11
LastEditors: 袁瑞 && 2502099390@qq.com
LastEditTime: 2026-06-12 15:01:10
FilePath: \sample_admin\backend\scripts\import_excel.py
Description: 
'''
from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import psycopg
from psycopg.types.json import Jsonb

from common import database_url, encrypt_text, normalize_text, sha256_text


CORE_FIELD_MAP = {
    "序": "sample_id",
    "条码号": "barcode_no",
    "实验号": "experiment_no",
    "病人号": "patient_no",
    "来源": "source_batch",
    "民族": "ethnicity",
    "性别": "sex",
    "年龄": "age",
    "类型": "visit_type",
    "科室": "department",
    "床号": "bed_no",
    "申请医生": "doctor",
    "诊断": "diagnosis",
    "申请项目": "request_items",
    "实验情况": "test_priority",
    "标本": "specimen_type",
    "设备": "device",
    "采集时间": "collection_time",
    "接收时间": "received_time",
    "上机时间": "machine_time",
    "审核时间": "review_time",
    "备注": "note",
}

PRIVATE_FIELDS = {"姓名", "身份证号", "电话", "地址"}
TIME_COLUMNS = {"collection_time", "received_time", "machine_time", "review_time"}
LAB_START_HEADER = "WBC"
SAMPLE_ID_PATTERN = re.compile(r"^[0-9]{6}$")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import sample Excel into PostgreSQL.")
    parser.add_argument("--file", required=True, help="Excel file path")
    parser.add_argument("--sheet", default="汇总数据", help="Worksheet name")
    parser.add_argument(
        "--allow-nonstandard-id",
        action="store_true",
        help="Allow sample_id values that are not 6 digit strings",
    )
    return parser.parse_args()


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_sample_id(value: Any) -> str | None:
    text = normalize_text(value)
    if text is None:
        return None
    return text.replace("\n", "").replace("\r", "").replace("\t", "").strip()


def parse_time(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    return value


def row_to_dict(headers: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
    data: dict[str, Any] = {}
    for header, value in zip(headers, values, strict=False):
        if not header:
            continue
        data[header] = json_value(value)
    return data


def build_payload(headers: list[str], row_values: tuple[Any, ...], row_number: int, allow_nonstandard: bool):
    raw_data = row_to_dict(headers, row_values)
    sample_id = clean_sample_id(raw_data.get("序"))
    if not sample_id:
        raise ValueError(f"第 {row_number} 行：序为空")
    if not allow_nonstandard and not SAMPLE_ID_PATTERN.match(sample_id):
        raise ValueError(f"第 {row_number} 行：序不是 6 位数字字符串：{sample_id}")

    core: dict[str, Any] = {"sample_id": sample_id}
    for excel_name, column_name in CORE_FIELD_MAP.items():
        if column_name == "sample_id":
            continue
        value = raw_data.get(excel_name)
        if column_name in TIME_COLUMNS:
            core[column_name] = parse_time(value)
        else:
            core[column_name] = normalize_text(value)

    private = {
        "name": normalize_text(raw_data.get("姓名")),
        "id_card_no": normalize_text(raw_data.get("身份证号")),
        "phone": normalize_text(raw_data.get("电话")),
        "address": normalize_text(raw_data.get("地址")),
    }

    lab_start_idx = headers.index(LAB_START_HEADER) if LAB_START_HEADER in headers else len(headers)
    lab_headers = set(headers[lab_start_idx:])

    visible_data: dict[str, Any] = {}
    lab_results: dict[str, Any] = {}
    ignored_fields: dict[str, Any] = {}
    known_core = set(CORE_FIELD_MAP)

    for header, value in raw_data.items():
        if header.startswith("None.") or header == "":
            ignored_fields[header] = value
        elif header in PRIVATE_FIELDS or header in known_core:
            continue
        elif header in lab_headers:
            lab_results[header] = value
        else:
            visible_data[header] = value

    return {
        "sample_id": sample_id,
        "core": core,
        "private": private,
        "raw_data": raw_data,
        "visible_data": visible_data,
        "lab_results": lab_results,
        "ignored_fields": ignored_fields,
    }


def import_excel(path: Path, sheet_name: str, allow_nonstandard: bool) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    headers = [normalize_text(cell.value) or "" for cell in ws[1]]

    with psycopg.connect(database_url()) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO sample_import_batches (
                    file_name, file_path, file_hash, total_rows, status, field_mapping
                )
                VALUES (%s, %s, %s, %s, 'imported', %s)
                RETURNING id
                """,
                (
                    path.name,
                    str(path),
                    file_hash(path),
                    max(ws.max_row - 1, 0),
                    Jsonb(CORE_FIELD_MAP),
                ),
            )
            import_batch_id = cur.fetchone()[0]

            success_rows = 0
            errors: list[str] = []
            seen: set[str] = set()

            for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                try:
                    payload = build_payload(headers, row, row_number, allow_nonstandard)
                    sample_id = payload["sample_id"]
                    if sample_id in seen:
                        raise ValueError(f"第 {row_number} 行：序在 Excel 内重复：{sample_id}")
                    seen.add(sample_id)
                    core = payload["core"]

                    cur.execute(
                        """
                        INSERT INTO samples (
                            sample_id, barcode_no, experiment_no, patient_no, source_batch,
                            ethnicity, sex, age, visit_type, department, bed_no, doctor,
                            diagnosis, request_items, test_priority, specimen_type, device,
                            collection_time, received_time, machine_time, review_time,
                            sample_status, note
                        )
                        VALUES (
                            %(sample_id)s, %(barcode_no)s, %(experiment_no)s, %(patient_no)s,
                            %(source_batch)s, %(ethnicity)s, %(sex)s, %(age)s, %(visit_type)s,
                            %(department)s, %(bed_no)s, %(doctor)s, %(diagnosis)s,
                            %(request_items)s, %(test_priority)s, %(specimen_type)s, %(device)s,
                            %(collection_time)s, %(received_time)s, %(machine_time)s,
                            %(review_time)s, 'in_storage', %(note)s
                        )
                        ON CONFLICT (sample_id) DO UPDATE SET
                            barcode_no = EXCLUDED.barcode_no,
                            experiment_no = EXCLUDED.experiment_no,
                            patient_no = EXCLUDED.patient_no,
                            source_batch = EXCLUDED.source_batch,
                            ethnicity = EXCLUDED.ethnicity,
                            sex = EXCLUDED.sex,
                            age = EXCLUDED.age,
                            visit_type = EXCLUDED.visit_type,
                            department = EXCLUDED.department,
                            bed_no = EXCLUDED.bed_no,
                            doctor = EXCLUDED.doctor,
                            diagnosis = EXCLUDED.diagnosis,
                            request_items = EXCLUDED.request_items,
                            test_priority = EXCLUDED.test_priority,
                            specimen_type = EXCLUDED.specimen_type,
                            device = EXCLUDED.device,
                            collection_time = EXCLUDED.collection_time,
                            received_time = EXCLUDED.received_time,
                            machine_time = EXCLUDED.machine_time,
                            review_time = EXCLUDED.review_time,
                            note = EXCLUDED.note
                        """,
                        core,
                    )

                    private = payload["private"]
                    cur.execute(
                        """
                        INSERT INTO participant_private_info (
                            sample_id, patient_no, name_encrypted, id_card_no_encrypted,
                            id_card_hash, phone_encrypted, phone_hash, address_encrypted
                        )
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                        ON CONFLICT (sample_id) DO UPDATE SET
                            patient_no = EXCLUDED.patient_no,
                            name_encrypted = EXCLUDED.name_encrypted,
                            id_card_no_encrypted = EXCLUDED.id_card_no_encrypted,
                            id_card_hash = EXCLUDED.id_card_hash,
                            phone_encrypted = EXCLUDED.phone_encrypted,
                            phone_hash = EXCLUDED.phone_hash,
                            address_encrypted = EXCLUDED.address_encrypted
                        """,
                        (
                            sample_id,
                            core.get("patient_no"),
                            encrypt_text(private["name"]),
                            encrypt_text(private["id_card_no"]),
                            sha256_text(private["id_card_no"]),
                            encrypt_text(private["phone"]),
                            sha256_text(private["phone"]),
                            encrypt_text(private["address"]),
                        ),
                    )

                    cur.execute(
                        """
                        INSERT INTO sample_raw_records (
                            sample_id, import_batch_id, source_file_name, sheet_name,
                            row_number, raw_data, visible_data, lab_results, ignored_fields
                        )
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                        ON CONFLICT (sample_id) DO UPDATE SET
                            import_batch_id = EXCLUDED.import_batch_id,
                            source_file_name = EXCLUDED.source_file_name,
                            sheet_name = EXCLUDED.sheet_name,
                            row_number = EXCLUDED.row_number,
                            raw_data = EXCLUDED.raw_data,
                            visible_data = EXCLUDED.visible_data,
                            lab_results = EXCLUDED.lab_results,
                            ignored_fields = EXCLUDED.ignored_fields
                        """,
                        (
                            sample_id,
                            import_batch_id,
                            path.name,
                            sheet_name,
                            row_number,
                            Jsonb(payload["raw_data"]),
                            Jsonb(payload["visible_data"]),
                            Jsonb(payload["lab_results"]),
                            Jsonb(payload["ignored_fields"]),
                        ),
                    )
                    success_rows += 1
                except Exception as exc:
                    errors.append(str(exc))

            cur.execute(
                """
                UPDATE sample_import_batches
                SET success_rows = %s,
                    failed_rows = %s,
                    error_report = %s,
                    updated_at = now()
                WHERE id = %s
                """,
                (success_rows, len(errors), Jsonb(errors), import_batch_id),
            )

        conn.commit()

    print(json.dumps({"success_rows": success_rows, "failed_rows": len(errors), "errors": errors[:20]}, ensure_ascii=False, indent=2))


def main() -> None:
    args = parse_args()
    import_excel(Path(args.file).resolve(), args.sheet, args.allow_nonstandard_id)


if __name__ == "__main__":
    main()
