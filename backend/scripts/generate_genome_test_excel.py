from pathlib import Path
import openpyxl

SOURCE_FILE = r"D:\desktop\sample_admin\附件\少数民族数据汇总 2025.5.1-31_filled.xlsx"
OUTPUT_FILE = Path(__file__).resolve().parents[1] / "uploads" / "genome_test_import.xlsx"


def main():
    wb_src = openpyxl.load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    ws = wb_src["汇总数据"]
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]

    sample_code_col = None
    for i, h in enumerate(headers):
        if h == "sample_code":
            sample_code_col = i
            break
    if sample_code_col is None:
        print("未找到 sample_code 列")
        wb_src.close()
        return

    codes = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = str(row[sample_code_col]).strip() if row[sample_code_col] else ""
        if code:
            codes.append(code)
        if len(codes) >= 10:
            break
    wb_src.close()

    if not codes:
        print("未提取到任何 sample_code")
        return

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "汇总数据"

    ws_out.append([
        "sample_code",
        "基因组数据状态",
        "数据类型",
        "测序公司",
        "测序平台",
        "测序仪器",
        "回库时间",
        "测序深度",
        "基因组QC",
        "最终状态",
        "不可用原因",
    ])

    for code in codes:
        ws_out.append([
            code,
            "已完成",
            "WGS",
            "华大基因",
            "DNBSEQ",
            "DNBSEQ-T7",
            "2026-06-15",
            "30x",
            "合格",
            "可用",
            "",
        ])

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(OUTPUT_FILE)
    print(f"已生成：{OUTPUT_FILE}")
    for c in codes:
        print(f"  {c}")


if __name__ == "__main__":
    main()
