from __future__ import annotations

import argparse
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SEQUENCE_HEADER = "序"
SEQUENCE_SOURCE_HEADERS = ("序", "序号", "条码号")
COLLECTION_TIME_HEADERS = ("采集时间", "采样时间", "采集日期", "采样日期")


def normalize(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def parse_date(value: str) -> date:
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(value.strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"无法解析日期：{value}")


def infer_date_from_filename(path: Path) -> date | None:
    match = re.search(r"(?<!\d)(20\d{2})[._-]?(\d{2})[._-]?(\d{2})(?!\d)", path.stem)
    if not match:
        return None
    try:
        return date(*(int(part) for part in match.groups()))
    except ValueError:
        return None


def find_header_column(worksheet: Any, headers: tuple[str, ...]) -> int | None:
    normalized = {
        normalize(cell.value): cell.column
        for cell in worksheet[1]
        if normalize(cell.value)
    }
    for header in headers:
        if header in normalized:
            return normalized[header]
    return None


def copy_cell_style(source: Any, target: Any) -> None:
    if source.has_style:
        target._style = copy(source._style)
    target.alignment = copy(source.alignment)


def prepare_workbook(
    input_path: Path,
    output_path: Path,
    sheet_name: str | None,
    sequence_source: str | None,
    collection_date: date | None,
) -> tuple[str, int, date | None]:
    keep_vba = input_path.suffix.lower() == ".xlsm"
    workbook = load_workbook(input_path, keep_vba=keep_vba)
    try:
        selected_sheet = sheet_name or (
            "汇总数据" if "汇总数据" in workbook.sheetnames else workbook.active.title
        )
        if selected_sheet not in workbook.sheetnames:
            raise ValueError(f"工作表不存在：{selected_sheet}")
        worksheet = workbook[selected_sheet]

        source_headers = (sequence_source,) if sequence_source else SEQUENCE_SOURCE_HEADERS
        sequence_column = find_header_column(worksheet, source_headers)
        if sequence_column is None:
            raise ValueError(f"未找到序号来源列，当前查找：{'、'.join(source_headers)}")

        if normalize(worksheet.cell(1, sequence_column).value) != SEQUENCE_HEADER:
            source_values = [worksheet.cell(row, sequence_column).value for row in range(1, worksheet.max_row + 1)]
            source_styles = [worksheet.cell(row, sequence_column) for row in range(1, worksheet.max_row + 1)]
            source_width = worksheet.column_dimensions[get_column_letter(sequence_column)].width
            worksheet.insert_cols(1)
            for row, (source_cell, value) in enumerate(zip(source_styles, source_values), start=1):
                target = worksheet.cell(row, 1)
                target.value = SEQUENCE_HEADER if row == 1 else value
                copy_cell_style(source_cell, target)
            worksheet.column_dimensions["A"].width = source_width or 12

        sequence_column = find_header_column(worksheet, (SEQUENCE_HEADER,))
        if sequence_column is None:
            raise ValueError("补充序列失败")

        collection_column = find_header_column(worksheet, COLLECTION_TIME_HEADERS)
        effective_date = collection_date or infer_date_from_filename(input_path)
        if collection_column is None:
            if effective_date is None:
                raise ValueError("未找到采集时间列，请使用 --collection-date 指定日期")
            collection_column = worksheet.max_column + 1
            worksheet.cell(1, collection_column).value = "采集时间"
            copy_cell_style(worksheet.cell(1, sequence_column), worksheet.cell(1, collection_column))
            worksheet.column_dimensions[get_column_letter(collection_column)].width = 20

        prepared_rows = 0
        seen_sequences: set[str] = set()
        errors: list[str] = []
        for row in range(2, worksheet.max_row + 1):
            sequence = normalize(worksheet.cell(row, sequence_column).value)
            if not sequence:
                continue
            if sequence in seen_sequences:
                errors.append(f"第 {row} 行：序重复：{sequence}")
                continue
            seen_sequences.add(sequence)

            collection_cell = worksheet.cell(row, collection_column)
            if not normalize(collection_cell.value):
                if effective_date is None:
                    errors.append(f"第 {row} 行：采集时间为空")
                    continue
                collection_cell.value = effective_date
                collection_cell.number_format = "yyyy-mm-dd"
            prepared_rows += 1

        if errors:
            preview = "；".join(errors[:20])
            suffix = f"；另有 {len(errors) - 20} 行" if len(errors) > 20 else ""
            raise ValueError(preview + suffix)
        if prepared_rows == 0:
            raise ValueError("没有找到可处理的样本行")

        workbook.save(output_path)
        return selected_sheet, prepared_rows, effective_date
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="补齐批量打码所需的序和采集时间列")
    parser.add_argument("--file", required=True, help="原始 .xlsx/.xlsm 文件")
    parser.add_argument("--sheet", help="工作表名称，默认优先汇总数据，否则使用活动工作表")
    parser.add_argument("--sequence-source", help="序号来源列，例如 条码号 或 序号")
    parser.add_argument("--collection-date", help="缺失采集时间时统一填写的日期，例如 2026-07-07")
    parser.add_argument("--output", help="输出文件，默认在原文件名后添加 _prepared")
    args = parser.parse_args()

    input_path = Path(args.file).expanduser().resolve()
    if input_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise SystemExit("仅支持 .xlsx / .xlsm 文件")
    if not input_path.exists():
        raise SystemExit(f"文件不存在：{input_path}")

    output_path = (
        Path(args.output).expanduser().resolve()
        if args.output
        else input_path.with_name(f"{input_path.stem}_prepared{input_path.suffix}")
    )
    default_date = parse_date(args.collection_date) if args.collection_date else None
    try:
        sheet, count, effective_date = prepare_workbook(
            input_path=input_path,
            output_path=output_path,
            sheet_name=args.sheet,
            sequence_source=args.sequence_source,
            collection_date=default_date,
        )
    except ValueError as exc:
        raise SystemExit(f"处理失败：{exc}") from exc

    date_message = effective_date.isoformat() if effective_date else "使用表内采集时间"
    print(f"处理完成：{output_path}")
    print(f"工作表：{sheet}；样本行：{count}；补充日期：{date_message}")


if __name__ == "__main__":
    main()
