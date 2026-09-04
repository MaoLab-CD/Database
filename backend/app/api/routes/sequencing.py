from __future__ import annotations

import logging
import shlex
import subprocess
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.deps import CurrentUser, get_current_user
from app.db.session import get_db
from app.services.sequencing_scanner import scan_sequencing_root

router = APIRouter()
logger = logging.getLogger(__name__)


SEQUENCING_EXPORT_HEADERS = [
    ("sample_code", "样本编码"),
    ("sample_id", "原始序号"),
    ("center_name", "样本中心"),
    ("center_code", "中心编码"),
    ("project_code", "项目编号"),
    ("sequencing_company", "测序公司"),
    ("sequencing_platform", "测序平台"),
    ("sequencing_instrument", "测序仪器"),
    ("sequencing_returned_at", "回库时间"),
    ("sequencing_depth", "测序深度"),
    ("genome_qc", "QC"),
    ("final_status", "最终状态"),
    ("missing_reason", "不可用原因"),
    ("raw_data_path", "RawData 路径"),
    ("r1_file_name", "R1 文件"),
    ("r2_file_name", "R2 文件"),
    ("md5_file_name", "MD5 文件"),
    ("total_size_bytes", "总大小（字节）"),
    ("data_status", "数据状态"),
    ("last_scanned_at", "最近扫描时间"),
]

SEQUENCING_DATA_STATUS_LABELS = {
    "unmatched": "未匹配",
    "matched": "已关联",
    "available": "数据完整",
    "complete": "数据完整",
    "incomplete": "数据不完整",
    "changed": "文件有变化",
    "missing": "文件缺失",
}


class SequencingRecordItem(BaseModel):
    id: int
    sample_id: str
    sample_code: str | None = None
    center_code: str | None = None
    center_name: str | None = None
    project_code: str
    raw_data_path: str
    r1_file_name: str | None = None
    r2_file_name: str | None = None
    md5_file_name: str | None = None
    total_size_bytes: int | None = None
    data_status: str
    scan_batch_id: int | None = None
    last_scanned_at: datetime
    file_modified_at: datetime | None = None
    sample_exists: bool
    genome_data_status: str | None = None
    data_type: str | None = None
    sequencing_company: str | None = None
    sequencing_platform: str | None = None
    sequencing_instrument: str | None = None
    sequencing_returned_at: datetime | None = None
    sequencing_depth: str | None = None
    genome_qc: str | None = None
    final_status: str | None = None
    missing_reason: str | None = None


class SequencingRecordListResponse(BaseModel):
    items: list[SequencingRecordItem]
    total: int
    total_size_bytes: int
    page: int
    page_size: int


class SequencingRecordDetail(BaseModel):
    record: dict[str, Any]
    scan_batch: dict[str, Any] | None = None


class ScanBatchItem(BaseModel):
    id: int
    scan_root_path: str
    scan_mode: str
    status: str
    started_at: datetime
    finished_at: datetime | None = None
    total_projects: int
    total_sample_dirs: int
    new_count: int
    existing_count: int
    changed_count: int
    missing_count: int
    unmatched_count: int
    incomplete_count: int
    error_message: str | None = None


class LocalScanRequest(BaseModel):
    root_path: str = "/data/sequencing"
    scan_mode: str = "incremental"
    projects: list[str] | None = None


class RemoteScanRequest(BaseModel):
    host: str
    port: int = 22
    username: str
    password: str | None = None
    backend_path: str = "/opt/sample-scan/backend"
    root_path: str = "/data/sequencing"
    scan_mode: str = "incremental"
    projects: list[str] | None = None
    python_command: str = "python"


class ScanTriggerResult(BaseModel):
    source: str
    status: str
    batch_id: int | None = None
    message: str
    counters: dict[str, int] | None = None
    output: str | None = None


class RemoteConnectionTestResult(BaseModel):
    ok: bool
    message: str
    output: str | None = None


def ensure_admin(current_user: CurrentUser) -> None:
    if current_user.role != "admin":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="只有管理员可以触发测序扫盘",
        )


def normalize_projects(projects: list[str] | None) -> list[str]:
    return [project.strip() for project in projects or [] if project.strip()]


def validate_scan_mode(scan_mode: str) -> str:
    if scan_mode not in {"full", "incremental"}:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="扫描模式无效")
    return scan_mode


def remote_python_command(payload: RemoteScanRequest) -> str:
    return (payload.python_command or "python").strip() or "python"


@dataclass(frozen=True)
class SshCommandResult:
    returncode: int
    stdout: str
    stderr: str


def run_ssh_command(payload: RemoteScanRequest, remote_command: str, timeout_seconds: int) -> SshCommandResult:
    if payload.password:
        try:
            import paramiko
        except ImportError:
            raise RuntimeError("当前后端环境未安装 paramiko，无法使用密码登录 SSH")

        client = paramiko.SSHClient()
        client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        try:
            client.connect(
                hostname=payload.host,
                port=payload.port,
                username=payload.username,
                password=payload.password,
                timeout=15,
                banner_timeout=15,
                auth_timeout=15,
                look_for_keys=False,
                allow_agent=False,
            )
            stdin, stdout, stderr = client.exec_command(remote_command, timeout=timeout_seconds)
            stdin.close()
            exit_status = stdout.channel.recv_exit_status()
            return SshCommandResult(
                returncode=exit_status,
                stdout=stdout.read().decode("utf-8", errors="replace"),
                stderr=stderr.read().decode("utf-8", errors="replace"),
            )
        finally:
            client.close()

    ssh_target = f"{payload.username}@{payload.host}"
    command = [
        "ssh",
        "-p",
        str(payload.port),
        "-o",
        "BatchMode=yes",
        "-o",
        "ConnectTimeout=15",
        ssh_target,
        remote_command,
    ]
    completed = subprocess.run(
        command,
        check=False,
        capture_output=True,
        text=True,
        timeout=timeout_seconds,
    )
    return SshCommandResult(
        returncode=completed.returncode,
        stdout=completed.stdout,
        stderr=completed.stderr,
    )


def build_record_filters(
    keyword: str | None,
    data_status: str | None,
    center_code: str | None,
) -> tuple[str, dict[str, Any]]:
    clauses = ["1 = 1"]
    params: dict[str, Any] = {}

    if keyword:
        clauses.append(
            """
            (
                r.sample_id ILIKE :keyword
                OR r.sample_code ILIKE :keyword
                OR r.project_code ILIKE :keyword
                OR r.raw_data_path ILIKE :keyword
                OR COALESCE(r.r1_file_name, '') ILIKE :keyword
                OR COALESCE(r.r2_file_name, '') ILIKE :keyword
            )
            """
        )
        params["keyword"] = f"%{keyword.strip()}%"

    if data_status:
        clauses.append("r.data_status = :data_status")
        params["data_status"] = data_status

    if center_code:
        clauses.append(
            """
            (
                r.sample_code LIKE :center_prefix
                OR EXISTS (
                    SELECT 1
                    FROM samples sf
                    WHERE sf.is_deleted = false
                      AND sf.center_code = :center_code
                      AND (
                        sf.sample_code = r.sample_code
                        OR sf.sample_code = r.sample_id
                        OR sf.sample_id = r.sample_id
                        OR sf.barcode_no = r.sample_id
                        OR (
                          r.sample_id ~ '^[0-9]{1,6}$'
                          AND lpad(COALESCE(sf.sample_id, ''), 6, '0') = lpad(r.sample_id, 6, '0')
                        )
                      )
                )
            )
            """
        )
        params["center_code"] = center_code
        params["center_prefix"] = f"{center_code}-%"

    return " AND ".join(clauses), params


@router.post("/scan/local", response_model=ScanTriggerResult)
def trigger_local_scan(
    payload: LocalScanRequest,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> ScanTriggerResult:
    ensure_admin(current_user)
    scan_mode = validate_scan_mode(payload.scan_mode)
    projects = normalize_projects(payload.projects)

    try:
        result = scan_sequencing_root(
            db=db,
            root=Path(payload.root_path).expanduser(),
            mode=scan_mode,
            projects=projects or None,
            created_by=current_user.id,
        )
    except FileNotFoundError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc))
    except NotADirectoryError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc))
    except Exception:
        logger.exception("本地测序目录扫描失败")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="本地扫描失败，请稍后重试或联系系统管理员",
        ) from None

    return ScanTriggerResult(
        source="local",
        status="completed",
        batch_id=result.batch_id,
        message="本地扫描完成",
        counters=result.counters,
    )


@router.post("/scan/remote", response_model=ScanTriggerResult)
def trigger_remote_scan(
    payload: RemoteScanRequest,
    current_user: CurrentUser = Depends(get_current_user),
) -> ScanTriggerResult:
    ensure_admin(current_user)
    scan_mode = validate_scan_mode(payload.scan_mode)
    projects = normalize_projects(payload.projects)
    if not payload.host.strip() or not payload.username.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="请填写 SSH 主机和用户名")
    if not 1 <= payload.port <= 65535:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="SSH 端口无效")

    python_command = remote_python_command(payload)
    remote_parts = [
        "cd",
        shlex.quote(payload.backend_path),
        "&&",
        python_command,
        "scripts/scan_sequencing.py",
        "--root",
        shlex.quote(payload.root_path),
        "--mode",
        shlex.quote(scan_mode),
    ]
    for project in projects:
        remote_parts.extend(["--project", shlex.quote(project)])
    remote_command = " ".join(remote_parts)

    try:
        completed = run_ssh_command(payload, remote_command, timeout_seconds=60 * 30)
    except FileNotFoundError:
        logger.exception("服务器缺少 SSH 命令")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="服务器缺少远程连接组件，请联系系统管理员",
        ) from None
    except RuntimeError:
        logger.exception("远程扫描组件不可用")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="远程连接组件不可用，请联系系统管理员",
        ) from None
    except subprocess.TimeoutExpired:
        logger.exception("远程扫描超时")
        raise HTTPException(
            status_code=status.HTTP_504_GATEWAY_TIMEOUT,
            detail="远程扫描超时，请检查网络连接后重试",
        ) from None
    except Exception:
        logger.exception("远程扫描执行失败")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="远程扫描失败，请检查连接配置后重试",
        ) from None

    output = "\n".join(part for part in [completed.stdout, completed.stderr] if part).strip()
    if completed.returncode != 0:
        logger.error("远程扫描返回失败状态：%s", output)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="远程扫描失败，请检查连接和扫描配置",
        )

    return ScanTriggerResult(
        source="remote",
        status="completed",
        message="远程扫描命令已完成，请刷新扫描批次查看入库结果",
        output=output,
    )


@router.post("/scan/remote/test", response_model=RemoteConnectionTestResult)
def test_remote_scan_connection(
    payload: RemoteScanRequest,
    current_user: CurrentUser = Depends(get_current_user),
) -> RemoteConnectionTestResult:
    ensure_admin(current_user)
    if not payload.host.strip() or not payload.username.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="请填写 SSH 主机和用户名")
    if not 1 <= payload.port <= 65535:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="SSH 端口无效")

    backend_path = payload.backend_path.rstrip("/")
    script_path = f"{backend_path}/scripts/scan_sequencing.py"
    python_command = remote_python_command(payload)
    remote_command = " ".join(
        [
            "if",
            "!",
            "test",
            "-d",
            shlex.quote(backend_path),
            ";",
            "then",
            "echo",
            shlex.quote(f"BACKEND_DIR_NOT_FOUND: {backend_path}"),
            ";",
            "exit",
            "2",
            ";",
            "fi",
            ";",
            "if",
            "!",
            "test",
            "-f",
            shlex.quote(script_path),
            ";",
            "then",
            "echo",
            shlex.quote(f"SCAN_SCRIPT_NOT_FOUND: {script_path}"),
            ";",
            "exit",
            "3",
            ";",
            "fi",
            ";",
            "if",
            "!",
            "(",
            python_command,
            "--version",
            ">/dev/null",
            "2>&1",
            ")",
            ";",
            "then",
            "echo",
            shlex.quote(f"PYTHON_COMMAND_NOT_FOUND: {python_command}"),
            ";",
            "exit",
            "4",
            ";",
            "fi",
            ";",
            "echo",
            shlex.quote("SSH_OK"),
        ]
    )
    try:
        completed = run_ssh_command(payload, remote_command, timeout_seconds=30)
    except FileNotFoundError:
        logger.exception("服务器缺少 SSH 命令")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="服务器缺少远程连接组件，请联系系统管理员",
        ) from None
    except RuntimeError:
        logger.exception("远程连接测试组件不可用")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="远程连接组件不可用，请联系系统管理员",
        ) from None
    except subprocess.TimeoutExpired:
        logger.exception("远程连接测试超时")
        raise HTTPException(
            status_code=status.HTTP_504_GATEWAY_TIMEOUT,
            detail="连接测试超时，请检查网络和连接配置",
        ) from None
    except Exception:
        logger.exception("远程连接测试失败")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="连接测试失败，请检查连接配置后重试",
        ) from None

    output = "\n".join(part for part in [completed.stdout, completed.stderr] if part).strip()
    if completed.returncode != 0:
        if "BACKEND_DIR_NOT_FOUND" in output:
            message = "SSH 登录成功，但远程脚本根目录不存在"
        elif "SCAN_SCRIPT_NOT_FOUND" in output:
            message = "SSH 登录成功，但远程扫盘脚本不存在"
        elif "PYTHON_COMMAND_NOT_FOUND" in output:
            message = "SSH 登录成功，但远程 Python 命令不可用"
        else:
            message = "SSH 可用性测试失败，请检查主机、端口、用户名、密码或远程目录"
        return RemoteConnectionTestResult(
            ok=False,
            message=message,
            output=None,
        )

    return RemoteConnectionTestResult(ok=True, message="SSH 连接和远程脚本检查通过", output=output)


@router.get("/records", response_model=SequencingRecordListResponse)
def list_sequencing_records(
    keyword: str | None = Query(default=None),
    data_status: str | None = Query(default=None),
    center_code: str | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> SequencingRecordListResponse:
    where_sql, params = build_record_filters(keyword, data_status, center_code)
    offset = (page - 1) * page_size

    summary = db.execute(
        text(
            f"""
            SELECT count(*) AS total,
                   COALESCE(sum(r.total_size_bytes), 0)::bigint AS total_size_bytes
            FROM sequencing_data_records r
            WHERE {where_sql}
            """
        ),
        params,
    ).mappings().one()

    rows = (
        db.execute(
            text(
                f"""
                SELECT r.id,
                       r.sample_id,
                       r.sample_code,
                       s.center_code,
                       c.center_name,
                       r.project_code,
                       r.raw_data_path,
                       r.r1_file_name,
                       r.r2_file_name,
                       r.md5_file_name,
                       r.total_size_bytes,
                       r.data_status,
                       r.scan_batch_id,
                       r.last_scanned_at,
                       r.file_modified_at,
                       (s.id IS NOT NULL) AS sample_exists,
                       g.genome_data_status,
                       g.data_type,
                       g.sequencing_company,
                       g.sequencing_platform,
                       g.sequencing_instrument,
                       g.sequencing_returned_at,
                       g.sequencing_depth,
                       g.genome_qc,
                       g.final_status,
                       g.missing_reason
                FROM sequencing_data_records r
                LEFT JOIN samples s
                  ON (
                    s.sample_code = r.sample_code
                    OR
                    s.sample_code = r.sample_id
                    OR s.sample_id = r.sample_id
                    OR s.barcode_no = r.sample_id
                    OR (
                      r.sample_id ~ '^[0-9]{{1,6}}$'
                      AND lpad(COALESCE(s.sample_id, ''), 6, '0') = lpad(r.sample_id, 6, '0')
                    )
                  )
                 AND s.is_deleted = false
                LEFT JOIN centers c ON c.center_code = s.center_code
                LEFT JOIN sample_genome_status g ON g.sample_code = r.sample_code
                WHERE {where_sql}
                ORDER BY r.last_scanned_at DESC, r.id DESC
                LIMIT :limit OFFSET :offset
                """
            ),
            {**params, "limit": page_size, "offset": offset},
        )
        .mappings()
        .all()
    )

    return SequencingRecordListResponse(
        items=[SequencingRecordItem(**row) for row in rows],
        total=int(summary["total"]),
        total_size_bytes=int(summary["total_size_bytes"]),
        page=page,
        page_size=page_size,
    )


def format_sequencing_export_value(key: str, value: Any) -> Any:
    if value is None:
        return ""
    if key == "data_status":
        return SEQUENCING_DATA_STATUS_LABELS.get(str(value), value)
    if isinstance(value, datetime):
        return value.strftime("%Y/%m/%d %H:%M")
    return value


@router.get("/records/export")
def export_sequencing_records(
    keyword: str | None = Query(default=None),
    data_status: str | None = Query(default=None),
    center_code: str | None = Query(default=None),
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> StreamingResponse:
    where_sql, params = build_record_filters(keyword, data_status, center_code)
    rows = (
        db.execute(
            text(
                f"""
                SELECT r.sample_id,
                       r.sample_code,
                       s.center_code,
                       c.center_name,
                       r.project_code,
                       r.raw_data_path,
                       r.r1_file_name,
                       r.r2_file_name,
                       r.md5_file_name,
                       r.total_size_bytes,
                       r.data_status,
                       r.last_scanned_at,
                       g.sequencing_company,
                       g.sequencing_platform,
                       g.sequencing_instrument,
                       g.sequencing_returned_at,
                       g.sequencing_depth,
                       g.genome_qc,
                       g.final_status,
                       g.missing_reason
                FROM sequencing_data_records r
                LEFT JOIN samples s
                  ON (
                    s.sample_code = r.sample_code
                    OR s.sample_code = r.sample_id
                    OR s.sample_id = r.sample_id
                    OR s.barcode_no = r.sample_id
                    OR (
                      r.sample_id ~ '^[0-9]{{1,6}}$'
                      AND lpad(COALESCE(s.sample_id, ''), 6, '0') = lpad(r.sample_id, 6, '0')
                    )
                  )
                 AND s.is_deleted = false
                LEFT JOIN centers c ON c.center_code = s.center_code
                LEFT JOIN sample_genome_status g ON g.sample_code = r.sample_code
                WHERE {where_sql}
                ORDER BY r.last_scanned_at DESC, r.id DESC
                """
            ),
            params,
        )
        .mappings()
        .all()
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "测序数据导出"
    worksheet.append([label for _, label in SEQUENCING_EXPORT_HEADERS])

    header_fill = PatternFill("solid", fgColor="EAF2FF")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="0F172A")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        worksheet.append([
            format_sequencing_export_value(key, row.get(key))
            for key, _ in SEQUENCING_EXPORT_HEADERS
        ])

    worksheet.freeze_panes = "A2"
    for column_index, (_, label) in enumerate(SEQUENCING_EXPORT_HEADERS, start=1):
        max_length = len(label)
        for cells in worksheet.iter_cols(
            min_col=column_index,
            max_col=column_index,
            min_row=2,
            max_row=worksheet.max_row,
        ):
            for cell in cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 10), 42)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    filename = f"测序数据导出_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )

@router.get("/records/{record_id}", response_model=SequencingRecordDetail)
def get_sequencing_record(
    record_id: int,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> SequencingRecordDetail:
    record = (
        db.execute(
            text(
                """
                SELECT id,
                       sample_id,
                       sample_code,
                       project_code,
                       data_root_path,
                       raw_data_path,
                       r1_file_name,
                       r2_file_name,
                       md5_file_name,
                       r1_file_size,
                       r2_file_size,
                       total_size_bytes,
                       file_modified_at,
                       md5_values,
                       data_status,
                       scan_batch_id,
                       last_scanned_at,
                       note,
                       created_at,
                       updated_at
                FROM sequencing_data_records
                WHERE id = :record_id
                """
            ),
            {"record_id": record_id},
        )
        .mappings()
        .first()
    )

    if record is None:
        return SequencingRecordDetail(record={})

    genome_status = None
    record_sample_code = record.get("sample_code")
    if record_sample_code:
        genome_row = (
            db.execute(
                text(
                    """
                    SELECT genome_data_status,
                           data_type,
                           sequencing_company,
                           sequencing_platform,
                           sequencing_instrument,
                           sequencing_returned_at,
                           sequencing_depth,
                           genome_qc,
                           final_status,
                           missing_reason
                    FROM sample_genome_status
                    WHERE sample_code = :sample_code
                    LIMIT 1
                    """
                ),
                {"sample_code": record_sample_code},
            )
            .mappings()
            .first()
        )
        if genome_row is not None:
            genome_status = dict(genome_row)

    scan_batch = None
    if record["scan_batch_id"]:
        scan_batch = (
            db.execute(
                text(
                    """
                    SELECT id,
                           scan_root_path,
                           scan_mode,
                           status,
                           started_at,
                           finished_at,
                           total_projects,
                           total_sample_dirs,
                           new_count,
                           existing_count,
                           changed_count,
                           missing_count,
                           unmatched_count,
                           incomplete_count,
                           error_message
                    FROM sequencing_scan_batches
                    WHERE id = :scan_batch_id
                    """
                ),
                {"scan_batch_id": record["scan_batch_id"]},
            )
            .mappings()
            .first()
        )

    return SequencingRecordDetail(
        record={**dict(record), "genome_status": genome_status},
        scan_batch=dict(scan_batch) if scan_batch else None,
    )


@router.get("/scan-batches", response_model=list[ScanBatchItem])
def list_scan_batches(
    limit: int = Query(default=10, ge=1, le=50),
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> list[ScanBatchItem]:
    rows = (
        db.execute(
            text(
                """
                SELECT id,
                       scan_root_path,
                       scan_mode,
                       status,
                       started_at,
                       finished_at,
                       total_projects,
                       total_sample_dirs,
                       new_count,
                       existing_count,
                       changed_count,
                       missing_count,
                       unmatched_count,
                       incomplete_count,
                       error_message
                FROM sequencing_scan_batches
                ORDER BY id DESC
                LIMIT :limit
                """
            ),
            {"limit": limit},
        )
        .mappings()
        .all()
    )

    return [ScanBatchItem(**row) for row in rows]
