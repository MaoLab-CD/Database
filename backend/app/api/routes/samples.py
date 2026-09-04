from __future__ import annotations

from io import BytesIO
from datetime import datetime
from decimal import Decimal
from typing import Any, Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.api.deps import CurrentUser, get_current_user
from app.core.security import decrypt_private_text
from app.db.session import get_db
from app.services.login_rate_limit import client_ip_from_request

router = APIRouter()


EXPORT_HEADERS = [
    ("sample_code", "样本编码"),
    ("sample_id", "原始序号"),
    ("center_name", "样本中心"),
    ("center_code", "中心编码"),
    ("barcode_no", "条码号"),
    ("experiment_no", "实验号"),
    ("patient_no", "病人号"),
    ("ethnicity", "民族"),
    ("sex", "性别"),
    ("age", "年龄"),
    ("visit_type", "类型"),
    ("department", "科室"),
    ("specimen_type", "样本类型"),
    ("sample_status", "样本状态"),
    ("sequencing_status", "数据状态"),
    ("storage_location", "冰箱位置"),
    ("collection_time", "采集时间"),
    ("review_time", "审核时间"),
    ("updated_at", "更新时间"),
]

SAMPLE_STATUS_LABELS = {
    "pending": "待确认",
    "not_stored": "未入库",
    "sequencing": "测序中",
    "in_storage": "在库",
    "checked_out": "已出库",
    "return_pending": "待归还复核",
    "consumed": "已用完",
    "lost": "丢失",
    "discarded": "废弃",
    "archived": "归档",
}

SEQUENCING_STATUS_LABELS = {
    "unmatched": "未匹配",
    "matched": "已关联",
    "complete": "数据完整",
    "incomplete": "数据不完整",
}


class SampleListItem(BaseModel):
    id: int
    sample_id: str
    sample_code: str | None = None
    center_code: str | None = None
    center_name: str | None = None
    sample_seq: str | None = None
    barcode_no: str | None = None
    experiment_no: str | None = None
    patient_no: str | None = None
    ethnicity: str | None = None
    sex: str | None = None
    age: str | None = None
    visit_type: str | None = None
    department: str | None = None
    specimen_type: str | None = None
    sample_status: str
    sequencing_status: str
    storage_location: str | None = None
    collection_time: datetime | None = None
    review_time: datetime | None = None
    updated_at: datetime


class SampleListResponse(BaseModel):
    items: list[SampleListItem]
    total: int
    page: int
    page_size: int


class SampleDetail(BaseModel):
    sample: dict[str, Any]
    lab_results: dict[str, Any]
    visible_data: dict[str, Any] = Field(default_factory=dict)
    source_file_name: str | None = None
    row_number: int | None = None
    genome_status: dict[str, Any] | None = None


class SamplePrivateInfo(BaseModel):
    id: int | None = None
    sample_id: str
    sample_code: str | None = None
    patient_no: str | None = None
    name: str | None = None
    id_card_no: str | None = None
    phone: str | None = None
    address: str | None = None
    accessed_at: datetime


class SampleCheckoutRequest(BaseModel):
    project_name: str | None = None
    purpose: str | None = None
    note: str | None = None


class SampleReturnRequest(BaseModel):
    used_amount: str | None = None
    unit: str | None = None
    used_volume_ul: Decimal | None = None
    purpose: str | None = None
    return_location: str | None = None
    note: str | None = None


class SampleOperationResult(BaseModel):
    sample_id: str
    sample_status: str
    related_record_id: int
    message: str


def ensure_sample_operator(current_user: CurrentUser) -> None:
    if current_user.role == "admin":
        return
    if current_user.permissions.get("sample_checkout"):
        return
    raise HTTPException(
        status_code=status.HTTP_403_FORBIDDEN,
        detail="当前账号没有扫码出入库权限",
    )


def build_filters(
    keyword: str | None,
    center_code: str | None,
    specimen_type: str | None,
    sample_status: str | None,
    sequencing_status: str | None,
) -> tuple[str, dict[str, Any]]:
    clauses = ["s.is_deleted = false"]
    params: dict[str, Any] = {}

    if keyword:
        clauses.append(
            """
            (
                s.sample_id ILIKE :keyword
                OR COALESCE(s.sample_code, '') ILIKE :keyword
                OR COALESCE(s.barcode_no, '') ILIKE :keyword
                OR COALESCE(s.patient_no, '') ILIKE :keyword
                OR COALESCE(s.experiment_no, '') ILIKE :keyword
                OR COALESCE(s.ethnicity, '') ILIKE :keyword
                OR COALESCE(s.visit_type, '') ILIKE :keyword
                OR COALESCE(s.department, '') ILIKE :keyword
                OR COALESCE(s.specimen_type, '') ILIKE :keyword
                OR COALESCE(c.center_name, '') ILIKE :keyword
            )
            """
        )
        params["keyword"] = f"%{keyword.strip()}%"

    if center_code:
        clauses.append("s.center_code = :center_code")
        params["center_code"] = center_code

    if specimen_type:
        clauses.append("s.specimen_type = :specimen_type")
        params["specimen_type"] = specimen_type

    if sample_status:
        clauses.append("s.sample_status = :sample_status")
        params["sample_status"] = sample_status

    if sequencing_status:
        clauses.append("s.sequencing_status = :sequencing_status")
        params["sequencing_status"] = sequencing_status

    return " AND ".join(clauses), params


@router.get("", response_model=SampleListResponse)
def list_samples(
    keyword: str | None = Query(default=None),
    center_code: str | None = Query(default=None),
    specimen_type: str | None = Query(default=None),
    sample_status: str | None = Query(default=None),
    sequencing_status: str | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> SampleListResponse:
    where_sql, params = build_filters(keyword, center_code, specimen_type, sample_status, sequencing_status)
    offset = (page - 1) * page_size

    total = db.execute(
        text(
            f"""
            SELECT count(*)
            FROM samples s
            LEFT JOIN centers c ON c.center_code = s.center_code
            WHERE {where_sql}
            """
        ),
        params,
    ).scalar_one()

    rows = (
        db.execute(
            text(
                f"""
                SELECT s.id,
                       s.sample_id,
                       s.sample_code,
                       s.center_code,
                       c.center_name,
                       s.sample_seq,
                       s.barcode_no,
                       s.experiment_no,
                       s.patient_no,
                       s.ethnicity,
                       s.sex,
                       s.age,
                       s.visit_type,
                       s.department,
                       s.specimen_type,
                       s.sample_status,
                       s.sequencing_status,
                       s.storage_location,
                       s.collection_time,
                       s.review_time,
                       s.updated_at
                FROM samples s
                LEFT JOIN centers c ON c.center_code = s.center_code
                LEFT JOIN specimen_types st ON st.name = s.specimen_type
                WHERE {where_sql}
                ORDER BY
                    regexp_replace(COALESCE(s.sample_code, s.sample_id), '[A-Za-z]$', ''),
                    COALESCE(st.sort_order, 9999),
                    COALESCE(s.sample_code, s.sample_id)
                LIMIT :limit OFFSET :offset
                """
            ),
            {**params, "limit": page_size, "offset": offset},
        )
        .mappings()
        .all()
    )

    return SampleListResponse(
        items=[SampleListItem(**row) for row in rows],
        total=int(total),
        page=page,
        page_size=page_size,
    )


def format_export_value(key: str, value: Any) -> Any:
    if value is None:
        return ""
    if key == "sample_status":
        return SAMPLE_STATUS_LABELS.get(str(value), value)
    if key == "sequencing_status":
        return SEQUENCING_STATUS_LABELS.get(str(value), value)
    if isinstance(value, datetime):
        return value.strftime("%Y/%m/%d %H:%M")
    return value


@router.get("/export")
def export_samples(
    keyword: str | None = Query(default=None),
    center_code: str | None = Query(default=None),
    specimen_type: str | None = Query(default=None),
    sample_status: str | None = Query(default=None),
    sequencing_status: str | None = Query(default=None),
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> StreamingResponse:
    where_sql, params = build_filters(keyword, center_code, specimen_type, sample_status, sequencing_status)
    rows = (
        db.execute(
            text(
                f"""
                SELECT s.id,
                       s.sample_id,
                       s.sample_code,
                       s.center_code,
                       c.center_name,
                       s.sample_seq,
                       s.barcode_no,
                       s.experiment_no,
                       s.patient_no,
                       s.ethnicity,
                       s.sex,
                       s.age,
                       s.visit_type,
                       s.department,
                       s.specimen_type,
                       s.sample_status,
                       s.sequencing_status,
                       s.storage_location,
                       s.collection_time,
                       s.review_time,
                       s.updated_at
                FROM samples s
                LEFT JOIN centers c ON c.center_code = s.center_code
                LEFT JOIN specimen_types st ON st.name = s.specimen_type
                WHERE {where_sql}
                ORDER BY
                    regexp_replace(COALESCE(s.sample_code, s.sample_id), '[A-Za-z]$', ''),
                    COALESCE(st.sort_order, 9999),
                    COALESCE(s.sample_code, s.sample_id)
                """
            ),
            params,
        )
        .mappings()
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "样本导出"
    ws.append([label for _, label in EXPORT_HEADERS])

    header_fill = PatternFill("solid", fgColor="EAF2FF")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="0F172A")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([format_export_value(key, row.get(key)) for key, _ in EXPORT_HEADERS])

    ws.freeze_panes = "A2"
    for column_index, (_, label) in enumerate(EXPORT_HEADERS, start=1):
        max_length = len(label)
        for cell in ws.iter_cols(
            min_col=column_index,
            max_col=column_index,
            min_row=2,
            max_row=ws.max_row,
        ):
            for item in cell:
                if item.value is not None:
                    max_length = max(max_length, len(str(item.value)))
        ws.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 10), 28)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"样本导出_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    encoded_filename = quote(filename)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.get("/{sample_id}", response_model=SampleDetail)
def get_sample_detail(
    sample_id: str,
    lookup_by: Literal["any", "sample_code"] = Query(default="any"),
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> SampleDetail:
    sample = (
        db.execute(
            text(
                """
                SELECT s.id,
                       s.sample_id,
                       s.sample_code,
                       s.center_code,
                       c.center_name,
                       s.sample_seq,
                       s.barcode_no,
                       s.experiment_no,
                       s.patient_no,
                       s.source_batch,
                       s.ethnicity,
                       s.sex,
                       s.age,
                       s.visit_type,
                       s.department,
                       s.bed_no,
                       s.doctor,
                       s.diagnosis,
                       s.request_items,
                       s.test_priority,
                       s.specimen_type,
                       s.device,
                       s.collection_time,
                       s.received_time,
                       s.machine_time,
                       s.review_time,
                       s.sample_status,
                       s.sequencing_status,
                       s.storage_location,
                       p.plate_code,
                       w.well_code,
                       pf.freezer_code AS plate_freezer_code,
                       pf.temperature_c AS plate_temperature_c,
                       p.layer_no AS plate_layer_no,
                       p.container_no AS plate_container_no,
                       p.location_note AS plate_location_note,
                       source.id AS source_sample_pk,
                       source.sample_id AS source_sample_id,
                       source.sample_code AS source_sample_code,
                       s.note,
                       s.created_at,
                       s.updated_at
                 FROM samples s
                 LEFT JOIN centers c ON c.center_code = s.center_code
                 LEFT JOIN sample_plate_wells w
                   ON w.sample_pk = s.id
                  AND w.removed_at IS NULL
                 LEFT JOIN sample_plates p ON p.id = w.plate_id
                 LEFT JOIN storage_freezers pf ON pf.id = p.freezer_id
                 LEFT JOIN sample_relations r
                   ON r.child_sample_pk = s.id
                  AND r.relation_type = 'dna_extraction'
                 LEFT JOIN samples source
                   ON source.id = r.parent_sample_pk
                  AND source.is_deleted = false
                 WHERE (
                    (
                        :sample_code_only = true
                        AND s.sample_code = :sample_id
                    )
                    OR (
                        :sample_code_only = false
                        AND (
                            s.id::text = :sample_id
                            OR s.sample_id = :sample_id
                            OR s.sample_code = :sample_id
                            OR s.barcode_no = :sample_id
                        )
                    )
                )
                  AND s.is_deleted = false
                """
            ),
            {
                "sample_id": sample_id,
                "sample_code_only": lookup_by == "sample_code",
            },
        )
        .mappings()
        .first()
    )

    raw_record = (
        db.execute(
            text(
                """
                SELECT source_file_name, row_number, visible_data, lab_results
                FROM sample_raw_records
                WHERE sample_pk = :sample_pk
                """
            ),
            {"sample_pk": sample["id"] if sample else None},
        )
        .mappings()
        .first()
    )

    if sample is None:
        return SampleDetail(sample={}, lab_results={})

    related_samples = [
        dict(row)
        for row in db.execute(
            text(
                """
                SELECT related.id,
                       related.sample_code,
                       related.sample_id,
                       related.specimen_type,
                       r.relation_type,
                       CASE
                           WHEN r.parent_sample_pk = :sample_pk THEN 'derived'
                           ELSE 'source'
                       END AS relation_direction
                FROM sample_relations r
                JOIN samples related
                  ON related.id = CASE
                       WHEN r.parent_sample_pk = :sample_pk THEN r.child_sample_pk
                       ELSE r.parent_sample_pk
                     END
                 AND related.is_deleted = false
                WHERE r.parent_sample_pk = :sample_pk
                   OR r.child_sample_pk = :sample_pk
                ORDER BY
                    CASE WHEN related.specimen_type = '全血' THEN 0 ELSE 1 END,
                    related.sample_code
                """
            ),
            {"sample_pk": sample["id"]},
        ).mappings()
    ]

    genome_status = None
    if sample.get("sample_code"):
        genome_row = (
            db.execute(
                text(
                    """
                    SELECT genome_data_status,
                           data_type,
                           sequencing_company,
                           sequencing_platform,
                           sequencing_instrument,
                           sequencing_returned_at,
                           sequencing_depth,
                           genome_qc,
                           final_status,
                           missing_reason
                    FROM sample_genome_status
                    WHERE sample_code = :sample_code
                    LIMIT 1
                    """
                ),
                {"sample_code": sample["sample_code"]},
            )
            .mappings()
            .first()
        )
        if genome_row is not None:
            genome_status = dict(genome_row)

    sample_data = dict(sample)
    sample_data["related_samples"] = related_samples
    return SampleDetail(
        sample=sample_data,
        lab_results=dict(raw_record["lab_results"] or {}) if raw_record else {},
        visible_data=dict(raw_record["visible_data"] or {}) if raw_record else {},
        source_file_name=raw_record["source_file_name"] if raw_record else None,
        row_number=raw_record["row_number"] if raw_record else None,
        genome_status=genome_status,
    )


@router.get("/{sample_id}/private-info", response_model=SamplePrivateInfo)
def get_sample_private_info(
    sample_id: str,
    request: Request,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> SamplePrivateInfo:
    if current_user.role != "admin":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="只有管理员可以查看敏感信息",
        )

    sample = (
        db.execute(
            text(
                """
                SELECT id, sample_id, sample_code
                FROM samples
                WHERE (
                    id::text = :sample_id
                    OR sample_id = :sample_id
                    OR sample_code = :sample_id
                    OR barcode_no = :sample_id
                )
                  AND is_deleted = false
                """
            ),
            {"sample_id": sample_id},
        )
        .mappings()
        .first()
    )

    if sample is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="样本不存在")

    row = (
        db.execute(
            text(
                """
                SELECT sample_id,
                       patient_no,
                       name_encrypted,
                       id_card_no_encrypted,
                       phone_encrypted,
                       address_encrypted
                FROM participant_private_info
                WHERE sample_pk = :sample_pk
                """
            ),
            {"sample_pk": sample["id"]},
        )
        .mappings()
        .first()
    )

    if row is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="该样本没有敏感信息记录",
        )

    viewed_fields = ["name", "id_card_no", "phone", "address"]
    db.execute(
        text(
            """
            INSERT INTO private_info_access_logs (
                sample_pk,
                sample_id,
                sample_code,
                viewer_id,
                viewer_role,
                viewed_fields,
                view_reason,
                access_type,
                ip_address,
                device_info
            )
            VALUES (
                :sample_pk,
                :sample_id,
                :sample_code,
                :viewer_id,
                :viewer_role,
                CAST(:viewed_fields AS jsonb),
                :view_reason,
                'view',
                :ip_address,
                :device_info
            )
            """
        ),
        {
            "sample_pk": sample["id"],
            "sample_id": sample["sample_id"],
            "sample_code": sample["sample_code"],
            "viewer_id": current_user.id,
            "viewer_role": current_user.role,
            "viewed_fields": '["name","id_card_no","phone","address"]',
            "view_reason": "admin_sample_detail_view",
            "ip_address": client_ip_from_request(request),
            "device_info": request.headers.get("user-agent"),
        },
    )
    db.commit()

    return SamplePrivateInfo(
        id=sample["id"],
        sample_id=row["sample_id"],
        sample_code=sample["sample_code"],
        patient_no=row["patient_no"],
        name=decrypt_private_text(row["name_encrypted"]),
        id_card_no=decrypt_private_text(row["id_card_no_encrypted"]),
        phone=decrypt_private_text(row["phone_encrypted"]),
        address=decrypt_private_text(row["address_encrypted"]),
        accessed_at=datetime.now(),
    )


@router.post("/{sample_id}/checkout", response_model=SampleOperationResult)
def checkout_sample(
    sample_id: str,
    request: Request,
    payload: SampleCheckoutRequest | None = None,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> SampleOperationResult:
    ensure_sample_operator(current_user)
    payload = payload or SampleCheckoutRequest()

    sample = (
        db.execute(
            text(
                """
                SELECT id, sample_id, sample_code, sample_status, storage_location
                FROM samples
                WHERE sample_code = :sample_id
                  AND is_deleted = false
                FOR UPDATE
                """
            ),
            {"sample_id": sample_id},
        )
        .mappings()
        .first()
    )

    if sample is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="正式样本编码不存在",
        )

    before_status = sample["sample_status"]
    if before_status != "in_storage":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"当前样本状态为 {before_status}，不可出库",
        )

    checkout_id = db.execute(
        text(
            """
            INSERT INTO sample_checkout_records (
                sample_pk,
                sample_id,
                checkout_user_id,
                project_name,
                purpose,
                status,
                note
            )
            VALUES (
                :sample_pk,
                :sample_id,
                :checkout_user_id,
                :project_name,
                :purpose,
                'active',
                :note
            )
            RETURNING id
            """
        ),
        {
            "sample_pk": sample["id"],
            "sample_id": sample["sample_id"],
            "checkout_user_id": current_user.id,
            "project_name": payload.project_name,
            "purpose": payload.purpose,
            "note": payload.note,
        },
    ).scalar_one()

    db.execute(
        text(
            """
            UPDATE samples
            SET sample_status = 'checked_out',
                current_holder_id = :current_holder_id,
                updated_by = :updated_by,
                updated_at = now()
            WHERE id = :sample_pk
            """
        ),
        {
            "sample_pk": sample["id"],
            "current_holder_id": current_user.id,
            "updated_by": current_user.id,
        },
    )

    db.execute(
        text(
            """
            INSERT INTO sample_movement_logs (
                sample_pk,
                sample_id,
                sample_code,
                action_type,
                before_status,
                after_status,
                before_location,
                operator_id,
                operator_role,
                related_record_id,
                detail,
                note,
                ip_address,
                device_info
            )
            VALUES (
                :sample_pk,
                :sample_id,
                :sample_code,
                'checkout',
                :before_status,
                'checked_out',
                :before_location,
                :operator_id,
                :operator_role,
                :related_record_id,
                CAST(:detail AS jsonb),
                :note,
                :ip_address,
                :device_info
            )
            """
        ),
        {
            "sample_pk": sample["id"],
            "sample_id": sample["sample_id"],
            "sample_code": sample["sample_code"],
            "before_status": before_status,
            "before_location": sample["storage_location"],
            "operator_id": current_user.id,
            "operator_role": current_user.role,
            "related_record_id": checkout_id,
            "detail": '{"source":"scan_workbench"}',
            "note": payload.note,
            "ip_address": client_ip_from_request(request),
            "device_info": request.headers.get("user-agent"),
        },
    )

    db.commit()

    return SampleOperationResult(
        sample_id=sample["sample_id"],
        sample_status="checked_out",
        related_record_id=checkout_id,
        message="样本已出库",
    )


@router.post("/{sample_id}/return-request", response_model=SampleOperationResult)
def request_sample_return(
    sample_id: str,
    request: Request,
    payload: SampleReturnRequest | None = None,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> SampleOperationResult:
    ensure_sample_operator(current_user)
    payload = payload or SampleReturnRequest()

    sample = (
        db.execute(
            text(
                """
                SELECT id, sample_id, sample_code, sample_status, storage_location
                FROM samples
                WHERE sample_code = :sample_id
                  AND is_deleted = false
                FOR UPDATE
                """
            ),
            {"sample_id": sample_id},
        )
        .mappings()
        .first()
    )

    if sample is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="正式样本编码不存在",
        )

    before_status = sample["sample_status"]
    if before_status != "checked_out":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"当前样本状态为 {before_status}，不可提交归还",
        )

    checkout = (
        db.execute(
            text(
                """
                SELECT id
                FROM sample_checkout_records
                WHERE sample_pk = :sample_pk AND status = 'active'
                ORDER BY checkout_time DESC
                LIMIT 1
                """
            ),
            {"sample_pk": sample["id"]},
        )
        .mappings()
        .first()
    )

    if checkout is None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="该样本没有 active 出库记录，无法提交归还",
        )

    return_id = db.execute(
        text(
            """
            INSERT INTO sample_return_records (
                sample_pk,
                sample_id,
                checkout_record_id,
                returned_by,
                purpose,
                used_amount,
                unit,
                used_volume_ul,
                return_location,
                status,
                note
            )
            VALUES (
                :sample_pk,
                :sample_id,
                :checkout_record_id,
                :returned_by,
                :purpose,
                :used_amount,
                :unit,
                :used_volume_ul,
                :return_location,
                'pending',
                :note
            )
            RETURNING id
            """
        ),
        {
            "sample_pk": sample["id"],
            "sample_id": sample["sample_id"],
            "checkout_record_id": checkout["id"],
            "returned_by": current_user.id,
            "purpose": payload.purpose,
            "used_amount": payload.used_amount
            or (str(payload.used_volume_ul) if payload.used_volume_ul is not None else None),
            "unit": payload.unit or ("ul" if payload.used_volume_ul is not None else None),
            "used_volume_ul": payload.used_volume_ul,
            "return_location": payload.return_location,
            "note": payload.note,
        },
    ).scalar_one()

    db.execute(
        text(
            """
            UPDATE samples
            SET sample_status = 'return_pending',
                updated_by = :updated_by,
                updated_at = now()
            WHERE id = :sample_pk
            """
        ),
        {"sample_pk": sample["id"], "updated_by": current_user.id},
    )

    db.execute(
        text(
            """
            INSERT INTO sample_movement_logs (
                sample_pk,
                sample_id,
                sample_code,
                action_type,
                before_status,
                after_status,
                before_location,
                after_location,
                operator_id,
                operator_role,
                related_record_id,
                detail,
                note,
                ip_address,
                device_info
            )
            VALUES (
                :sample_pk,
                :sample_id,
                :sample_code,
                'return_request',
                :before_status,
                'return_pending',
                :before_location,
                :after_location,
                :operator_id,
                :operator_role,
                :related_record_id,
                CAST(:detail AS jsonb),
                :note,
                :ip_address,
                :device_info
            )
            """
        ),
        {
            "sample_pk": sample["id"],
            "sample_id": sample["sample_id"],
            "sample_code": sample["sample_code"],
            "before_status": before_status,
            "before_location": sample["storage_location"],
            "after_location": payload.return_location,
            "operator_id": current_user.id,
            "operator_role": current_user.role,
            "related_record_id": return_id,
            "detail": '{"source":"scan_workbench"}',
            "note": payload.note,
            "ip_address": client_ip_from_request(request),
            "device_info": request.headers.get("user-agent"),
        },
    )

    db.commit()

    return SampleOperationResult(
        sample_id=sample["sample_id"],
        sample_status="return_pending",
        related_record_id=return_id,
        message="已提交归还申请，等待管理员复核",
    )
