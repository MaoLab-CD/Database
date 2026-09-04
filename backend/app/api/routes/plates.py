from __future__ import annotations

import json
from datetime import datetime
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.deps import CurrentUser, get_current_user
from app.db.session import get_db
from app.services.login_rate_limit import client_ip_from_request


router = APIRouter()

PLATE_EXPORT_HEADERS = [
    ("plate_code", "板号"),
    ("well_code", "孔位"),
    ("sample_code", "DNA 样本编码"),
    ("sample_id", "原始序号"),
    ("barcode_no", "条码号"),
    ("experiment_no", "实验号"),
    ("patient_no", "病人号"),
    ("center_name", "样本中心"),
    ("center_code", "中心编码"),
    ("ethnicity", "民族"),
    ("sex", "性别"),
    ("age", "年龄"),
    ("department", "科室"),
    ("specimen_type", "样本类型"),
    ("sample_status", "样本状态"),
    ("sequencing_status", "数据状态"),
    ("storage_location", "样本存储位置"),
    ("freezer_code", "冰箱号"),
    ("temperature_c", "温度（°C）"),
    ("layer_no", "层"),
    ("container_no", "板架号"),
    ("location_note", "板位备注"),
    ("source_sample_code", "来源血样编码"),
    ("source_sample_id", "来源血样原始序号"),
    ("placed_at", "放入时间"),
    ("updated_at", "样本更新时间"),
]

SAMPLE_STATUS_LABELS = {
    "pending": "待确认",
    "not_stored": "未入库",
    "sequencing": "测序中",
    "in_storage": "在库",
    "checked_out": "已出库",
    "return_pending": "待归还复核",
    "consumed": "已用完",
    "lost": "丢失",
    "discarded": "废弃",
    "archived": "归档",
}

SEQUENCING_STATUS_LABELS = {
    "none": "无数据",
    "available": "可用",
    "incomplete": "数据不完整",
    "unmatched": "未匹配",
    "changed": "文件有变化",
    "missing": "文件缺失",
    "archived": "归档",
}


class PlateListItem(BaseModel):
    id: int
    plate_code: str
    plate_type: str
    row_count: int
    column_count: int
    freezer_code: str | None = None
    temperature_c: float | None = None
    layer_no: str | None = None
    container_no: str | None = None
    location_note: str | None = None
    is_active: bool
    occupied_count: int
    in_storage_count: int
    checked_out_count: int
    return_pending_count: int
    unavailable_count: int
    updated_at: datetime


class PlateListResponse(BaseModel):
    items: list[PlateListItem]
    total: int
    page: int
    page_size: int


class PlateWellItem(BaseModel):
    well_code: str
    sample_pk: int
    sample_id: str
    sample_code: str | None = None
    specimen_type: str | None = None
    sample_status: str
    source_sample_pk: int | None = None
    source_sample_id: str | None = None
    source_sample_code: str | None = None
    placed_at: datetime


class PlateDetailResponse(BaseModel):
    plate: PlateListItem
    wells: list[PlateWellItem]


class PlateCheckoutRequest(BaseModel):
    well_codes: list[str] = Field(min_length=1, max_length=96)
    project_name: str | None = None
    purpose: str | None = None
    note: str | None = None


class PlateCheckoutItem(BaseModel):
    well_code: str
    sample_id: str
    sample_code: str | None = None
    checkout_record_id: int


class PlateCheckoutResponse(BaseModel):
    plate_code: str
    checkout_count: int
    items: list[PlateCheckoutItem]
    message: str


def ensure_sample_operator(current_user: CurrentUser) -> None:
    if current_user.role == "admin" or current_user.permissions.get("sample_checkout"):
        return
    raise HTTPException(
        status_code=status.HTTP_403_FORBIDDEN,
        detail="当前账号没有扫码出入库权限",
    )


PLATE_SELECT = """
    SELECT p.id,
           p.plate_code,
           p.plate_type,
           p.row_count,
           p.column_count,
           f.freezer_code,
           f.temperature_c,
           p.layer_no,
           p.container_no,
           p.location_note,
           p.is_active,
           count(w.id) FILTER (
               WHERE s.id IS NOT NULL AND s.is_deleted = false
           )::int AS occupied_count,
           count(w.id) FILTER (
               WHERE s.is_deleted = false AND s.sample_status = 'in_storage'
           )::int AS in_storage_count,
           count(w.id) FILTER (
               WHERE s.is_deleted = false AND s.sample_status = 'checked_out'
           )::int AS checked_out_count,
           count(w.id) FILTER (
               WHERE s.is_deleted = false AND s.sample_status = 'return_pending'
           )::int AS return_pending_count,
           count(w.id) FILTER (
               WHERE s.is_deleted = false
                 AND s.sample_status IN ('consumed', 'lost', 'discarded', 'archived')
           )::int AS unavailable_count,
           p.updated_at
    FROM sample_plates p
    LEFT JOIN storage_freezers f ON f.id = p.freezer_id
    LEFT JOIN sample_plate_wells w
      ON w.plate_id = p.id
     AND w.removed_at IS NULL
    LEFT JOIN samples s ON s.id = w.sample_pk
"""


def plate_group_by() -> str:
    return """
        GROUP BY p.id,
                 p.plate_code,
                 p.plate_type,
                 p.row_count,
                 p.column_count,
                 f.freezer_code,
                 f.temperature_c,
                 p.layer_no,
                 p.container_no,
                 p.location_note,
                 p.is_active,
                 p.updated_at
    """


def format_plate_export_value(key: str, value: object) -> object:
    if value is None:
        return ""
    if key == "sample_status":
        return SAMPLE_STATUS_LABELS.get(str(value), value)
    if key == "sequencing_status":
        return SEQUENCING_STATUS_LABELS.get(str(value), value)
    if isinstance(value, datetime):
        return value.strftime("%Y/%m/%d %H:%M")
    return value


@router.get("", response_model=PlateListResponse)
def list_plates(
    keyword: str | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> PlateListResponse:
    del current_user
    where_sql = "1 = 1"
    params: dict[str, object] = {}
    if keyword and keyword.strip():
        where_sql = """
            (
                p.plate_code ILIKE :keyword
                OR COALESCE(f.freezer_code, '') ILIKE :keyword
                OR COALESCE(p.layer_no, '') ILIKE :keyword
                OR COALESCE(p.container_no, '') ILIKE :keyword
            )
        """
        params["keyword"] = f"%{keyword.strip()}%"

    total = int(
        db.execute(
            text(
                f"""
                SELECT count(*)
                FROM sample_plates p
                LEFT JOIN storage_freezers f ON f.id = p.freezer_id
                WHERE {where_sql}
                """
            ),
            params,
        ).scalar_one()
    )

    params.update({"limit": page_size, "offset": (page - 1) * page_size})
    rows = (
        db.execute(
            text(
                f"""
                {PLATE_SELECT}
                WHERE {where_sql}
                {plate_group_by()}
                ORDER BY p.plate_code
                LIMIT :limit OFFSET :offset
                """
            ),
            params,
        )
        .mappings()
        .all()
    )
    return PlateListResponse(
        items=[PlateListItem(**dict(row)) for row in rows],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get("/{plate_code}/export")
def export_plate_samples(
    plate_code: str,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> StreamingResponse:
    del current_user
    plate = (
        db.execute(
            text(
                """
                SELECT p.id,
                       p.plate_code,
                       f.freezer_code,
                       f.temperature_c,
                       p.layer_no,
                       p.container_no,
                       p.location_note
                FROM sample_plates p
                LEFT JOIN storage_freezers f ON f.id = p.freezer_id
                WHERE p.plate_code = :plate_code
                """
            ),
            {"plate_code": plate_code},
        )
        .mappings()
        .first()
    )
    if plate is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="DNA 板不存在")

    rows = (
        db.execute(
            text(
                """
                SELECT p.plate_code,
                       w.well_code,
                       s.sample_code,
                       s.sample_id,
                       s.barcode_no,
                       s.experiment_no,
                       s.patient_no,
                       c.center_name,
                       s.center_code,
                       s.ethnicity,
                       s.sex,
                       s.age,
                       s.department,
                       s.specimen_type,
                       s.sample_status,
                       s.sequencing_status,
                       s.storage_location,
                       f.freezer_code,
                       f.temperature_c,
                       p.layer_no,
                       p.container_no,
                       p.location_note,
                       source.sample_code AS source_sample_code,
                       source.sample_id AS source_sample_id,
                       w.placed_at,
                       s.updated_at
                FROM sample_plate_wells w
                JOIN sample_plates p ON p.id = w.plate_id
                JOIN samples s
                  ON s.id = w.sample_pk
                 AND s.is_deleted = false
                LEFT JOIN centers c ON c.center_code = s.center_code
                LEFT JOIN storage_freezers f ON f.id = p.freezer_id
                LEFT JOIN sample_relations r
                  ON r.child_sample_pk = s.id
                 AND r.relation_type = 'dna_extraction'
                LEFT JOIN samples source
                  ON source.id = r.parent_sample_pk
                 AND source.is_deleted = false
                WHERE w.plate_id = :plate_id
                  AND w.removed_at IS NULL
                ORDER BY substring(w.well_code FROM 1 FOR 1),
                         substring(w.well_code FROM 2)::int
                """
            ),
            {"plate_id": plate["id"]},
        )
        .mappings()
        .all()
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "DNA板样本"
    worksheet.append([label for _, label in PLATE_EXPORT_HEADERS])

    header_fill = PatternFill("solid", fgColor="EAF2FF")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="0F172A")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        worksheet.append(
            [
                format_plate_export_value(key, row.get(key))
                for key, _ in PLATE_EXPORT_HEADERS
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_index, (_, label) in enumerate(PLATE_EXPORT_HEADERS, start=1):
        max_length = len(label)
        if worksheet.max_row >= 2:
            for column in worksheet.iter_cols(
                min_col=column_index,
                max_col=column_index,
                min_row=2,
                max_row=worksheet.max_row,
            ):
                for cell in column:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_length + 2, 10),
            30,
        )

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    filename = f"DNA板_{plate_code}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
        },
    )


@router.get("/{plate_code}", response_model=PlateDetailResponse)
def get_plate_detail(
    plate_code: str,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> PlateDetailResponse:
    del current_user
    plate_row = (
        db.execute(
            text(
                f"""
                {PLATE_SELECT}
                WHERE p.plate_code = :plate_code
                {plate_group_by()}
                """
            ),
            {"plate_code": plate_code},
        )
        .mappings()
        .first()
    )
    if plate_row is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="DNA 板不存在")

    well_rows = (
        db.execute(
            text(
                """
                SELECT w.well_code,
                       s.id AS sample_pk,
                       s.sample_id,
                       s.sample_code,
                       s.specimen_type,
                       s.sample_status,
                       source.id AS source_sample_pk,
                       source.sample_id AS source_sample_id,
                       source.sample_code AS source_sample_code,
                       w.placed_at
                FROM sample_plate_wells w
                JOIN samples s
                  ON s.id = w.sample_pk
                 AND s.is_deleted = false
                LEFT JOIN sample_relations r
                  ON r.child_sample_pk = s.id
                 AND r.relation_type = 'dna_extraction'
                LEFT JOIN samples source
                  ON source.id = r.parent_sample_pk
                 AND source.is_deleted = false
                WHERE w.plate_id = :plate_id
                  AND w.removed_at IS NULL
                ORDER BY substring(w.well_code FROM 1 FOR 1),
                         substring(w.well_code FROM 2)::int
                """
            ),
            {"plate_id": plate_row["id"]},
        )
        .mappings()
        .all()
    )

    return PlateDetailResponse(
        plate=PlateListItem(**dict(plate_row)),
        wells=[PlateWellItem(**dict(row)) for row in well_rows],
    )


@router.post("/{plate_code}/checkout", response_model=PlateCheckoutResponse)
def checkout_plate_wells(
    plate_code: str,
    payload: PlateCheckoutRequest,
    request: Request,
    db: Session = Depends(get_db),
    current_user: CurrentUser = Depends(get_current_user),
) -> PlateCheckoutResponse:
    ensure_sample_operator(current_user)
    well_codes = [code.strip().upper() for code in payload.well_codes if code.strip()]
    if not well_codes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="请选择要出库的 DNA 孔位")
    if len(well_codes) != len(set(well_codes)):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="出库孔位不能重复")

    plate = (
        db.execute(
            text(
                """
                SELECT id, plate_code
                FROM sample_plates
                WHERE plate_code = :plate_code
                  AND is_active = true
                FOR UPDATE
                """
            ),
            {"plate_code": plate_code},
        )
        .mappings()
        .first()
    )
    if plate is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="DNA 板不存在或已停用")

    samples = (
        db.execute(
            text(
                """
                SELECT w.well_code,
                       s.id,
                       s.sample_id,
                       s.sample_code,
                       s.specimen_type,
                       s.sample_status,
                       s.storage_location
                FROM sample_plate_wells w
                JOIN samples s
                  ON s.id = w.sample_pk
                 AND s.is_deleted = false
                WHERE w.plate_id = :plate_id
                  AND w.removed_at IS NULL
                  AND w.well_code = ANY(:well_codes)
                ORDER BY substring(w.well_code FROM 1 FOR 1),
                         substring(w.well_code FROM 2)::int
                FOR UPDATE OF s
                """
            ),
            {"plate_id": plate["id"], "well_codes": well_codes},
        )
        .mappings()
        .all()
    )

    found_codes = {row["well_code"] for row in samples}
    missing_codes = [code for code in well_codes if code not in found_codes]
    if missing_codes:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"以下孔位为空或样本已不存在：{'、'.join(missing_codes)}",
        )

    non_dna = [row["well_code"] for row in samples if row["specimen_type"] != "DNA"]
    if non_dna:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"以下孔位不是 DNA 样本：{'、'.join(non_dna)}",
        )

    unavailable = [
        f"{row['well_code']}（{row['sample_status']}）"
        for row in samples
        if row["sample_status"] != "in_storage"
    ]
    if unavailable:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"以下 DNA 当前不可出库：{'、'.join(unavailable)}",
        )

    result_items: list[PlateCheckoutItem] = []
    for sample in samples:
        checkout_id = db.execute(
            text(
                """
                INSERT INTO sample_checkout_records (
                    sample_pk,
                    sample_id,
                    checkout_user_id,
                    project_name,
                    purpose,
                    status,
                    note
                )
                VALUES (
                    :sample_pk,
                    :sample_id,
                    :checkout_user_id,
                    :project_name,
                    :purpose,
                    'active',
                    :note
                )
                RETURNING id
                """
            ),
            {
                "sample_pk": sample["id"],
                "sample_id": sample["sample_id"],
                "checkout_user_id": current_user.id,
                "project_name": payload.project_name,
                "purpose": payload.purpose,
                "note": payload.note,
            },
        ).scalar_one()

        db.execute(
            text(
                """
                UPDATE samples
                SET sample_status = 'checked_out',
                    current_holder_id = :current_holder_id,
                    updated_by = :updated_by,
                    updated_at = now()
                WHERE id = :sample_pk
                """
            ),
            {
                "sample_pk": sample["id"],
                "current_holder_id": current_user.id,
                "updated_by": current_user.id,
            },
        )

        db.execute(
            text(
                """
                INSERT INTO sample_movement_logs (
                    sample_pk,
                    sample_id,
                    sample_code,
                    action_type,
                    before_status,
                    after_status,
                    before_location,
                    operator_id,
                    operator_role,
                    related_record_id,
                    detail,
                    note,
                    ip_address,
                    device_info
                )
                VALUES (
                    :sample_pk,
                    :sample_id,
                    :sample_code,
                    'checkout',
                    'in_storage',
                    'checked_out',
                    :before_location,
                    :operator_id,
                    :operator_role,
                    :related_record_id,
                    CAST(:detail AS jsonb),
                    :note,
                    :ip_address,
                    :device_info
                )
                """
            ),
            {
                "sample_pk": sample["id"],
                "sample_id": sample["sample_id"],
                "sample_code": sample["sample_code"],
                "before_location": sample["storage_location"],
                "operator_id": current_user.id,
                "operator_role": current_user.role,
                "related_record_id": checkout_id,
                "detail": json.dumps(
                    {
                        "source": "plate_scan_workbench",
                        "plate_code": plate["plate_code"],
                        "well_code": sample["well_code"],
                    },
                    ensure_ascii=False,
                ),
                "note": payload.note,
                "ip_address": client_ip_from_request(request),
                "device_info": request.headers.get("user-agent"),
            },
        )
        result_items.append(
            PlateCheckoutItem(
                well_code=sample["well_code"],
                sample_id=sample["sample_id"],
                sample_code=sample["sample_code"],
                checkout_record_id=checkout_id,
            )
        )

    db.commit()
    return PlateCheckoutResponse(
        plate_code=plate["plate_code"],
        checkout_count=len(result_items),
        items=result_items,
        message=f"已出库 {len(result_items)} 份 DNA 样本",
    )
