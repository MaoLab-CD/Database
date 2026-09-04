from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
import xlrd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


LEGACY_XLS_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
OOXML_ZIP_MAGICS = (
    b"PK\x03\x04",
    b"PK\x05\x06",
    b"PK\x07\x08",
)


def is_legacy_xls(content: bytes, suffix: str | None = None) -> bool:
    if content.startswith(LEGACY_XLS_MAGIC):
        return True
    if content.startswith(OOXML_ZIP_MAGICS):
        return False
    return (suffix or "").lower() == ".xls"


def _cell_value(book: Any, cell: Any) -> Any:
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, f"#ERROR_{cell.value}")
    if isinstance(cell.value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", cell.value)
    return cell.value


def convert_xls_to_xlsx(content: bytes) -> bytes:
    try:
        source = xlrd.open_workbook(file_contents=content, formatting_info=True)
    except Exception as exc:
        raise ValueError(f"无法读取旧版 .xls 文件：{exc}") from exc

    target = Workbook()
    target.remove(target.active)
    try:
        used_titles: set[str] = set()
        for sheet_index in range(source.nsheets):
            source_sheet = source.sheet_by_index(sheet_index)
            base_title = source_sheet.name[:31] or f"Sheet{sheet_index + 1}"
            title = base_title
            suffix_number = 2
            while title in used_titles:
                suffix_text = f"_{suffix_number}"
                title = f"{base_title[:31 - len(suffix_text)]}{suffix_text}"
                suffix_number += 1
            used_titles.add(title)
            target_sheet = target.create_sheet(title)

            for row_index in range(source_sheet.nrows):
                for column_index in range(source_sheet.ncols):
                    source_cell = source_sheet.cell(row_index, column_index)
                    target_cell = target_sheet.cell(row_index + 1, column_index + 1)
                    target_cell.value = _cell_value(source, source_cell)
                    if source_cell.xf_index < len(source.xf_list):
                        xf = source.xf_list[source_cell.xf_index]
                        number_format = source.format_map.get(xf.format_key)
                        if number_format and number_format.format_str:
                            target_cell.number_format = number_format.format_str

            for row_low, row_high, column_low, column_high in source_sheet.merged_cells:
                if row_high > row_low and column_high > column_low:
                    target_sheet.merge_cells(
                        start_row=row_low + 1,
                        end_row=row_high,
                        start_column=column_low + 1,
                        end_column=column_high,
                    )

            for column_index, column_info in source_sheet.colinfo_map.items():
                if column_info.width:
                    column_letter = openpyxl.utils.get_column_letter(column_index + 1)
                    target_sheet.column_dimensions[column_letter].width = column_info.width / 256
                target_sheet.column_dimensions[
                    openpyxl.utils.get_column_letter(column_index + 1)
                ].hidden = bool(column_info.hidden)

            for row_index, row_info in source_sheet.rowinfo_map.items():
                if row_info.height:
                    target_sheet.row_dimensions[row_index + 1].height = row_info.height / 20
                target_sheet.row_dimensions[row_index + 1].hidden = bool(row_info.hidden)

        output = BytesIO()
        target.save(output)
        return output.getvalue()
    finally:
        target.close()
        source.release_resources()


def normalize_excel_content(content: bytes, suffix: str | None = None) -> tuple[bytes, str]:
    if is_legacy_xls(content, suffix):
        return convert_xls_to_xlsx(content), ".xlsx"
    return content, (suffix or ".xlsx").lower()


def load_excel_workbook(
    content: bytes,
    suffix: str | None = None,
    *,
    read_only: bool = False,
    data_only: bool = False,
    keep_vba: bool = False,
) -> Any:
    normalized, normalized_suffix = normalize_excel_content(content, suffix)
    return openpyxl.load_workbook(
        BytesIO(normalized),
        read_only=read_only,
        data_only=data_only,
        keep_vba=keep_vba and normalized_suffix == ".xlsm",
    )


def load_excel_path(
    path: Path,
    *,
    read_only: bool = False,
    data_only: bool = False,
    keep_vba: bool = False,
) -> Any:
    return load_excel_workbook(
        path.read_bytes(),
        path.suffix.lower(),
        read_only=read_only,
        data_only=data_only,
        keep_vba=keep_vba,
    )
