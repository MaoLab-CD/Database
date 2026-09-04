from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any
from zipfile import BadZipFile, ZipFile, is_zipfile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.services.excel_importer import (
    CSV_SHEET_NAME,
    COLLECTION_TIME_HEADERS,
    CORE_FIELD_MAP,
    SAMPLE_ID_HEADERS,
    SPECIMEN_TYPE_HEADERS,
    build_payload,
    parse_sample_code,
    read_import_rows,
)
from app.services.sample_code_excel import (
    SampleCodeAssignment,
    SampleCodeInput,
    allocate_sample_codes,
    parse_collection_date,
)
from app.services.specimen_types import get_active_specimen_type_names
from app.services.workbook_loader import load_excel_path

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_WORKSHEETS = 20
MAX_ROWS = 20_000
MAX_COLUMNS = 300
MAX_STORED_ERRORS = 500
TEMPLATE_VERSION = "V1"
ALLOWED_HOSPITAL_SUFFIXES = {".xls", ".xlsx", ".xlsm", ".csv"}
REVIEW_SAMPLE_CODE_HEADER = "sample_code"
REVIEW_PREVIEW_RESULT_HEADERS = ("校验结果", "校验说明")
SENSITIVE_REVIEW_HEADERS = {
    "姓名",
    "患者姓名",
    "身份证号",
    "身份证号码",
    "证件号码",
    "电话",
    "手机号",
    "手机号码",
    "联系电话",
    "地址",
    "家庭住址",
    "居住地址",
}

FORBIDDEN_PACKAGE_PARTS = (
    "vbaproject.bin",
    "externallinks/",
    "embeddings/",
    "oleobjects/",
    "activex/",
)


@dataclass(frozen=True)
class HospitalValidationResult:
    total_rows: int
    valid_rows: int
    error_rows: int
    new_rows: int
    duplicate_rows: int
    errors: list[dict[str, Any]]
    summary: dict[str, Any]

    @property
    def is_valid(self) -> bool:
        return self.total_rows > 0 and self.error_rows == 0


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file_obj:
        for chunk in iter(lambda: file_obj.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def inspect_hospital_upload_file(path: Path) -> None:
    size = path.stat().st_size
    if size <= 0:
        raise ValueError("上传文件为空")
    if size > MAX_UPLOAD_BYTES:
        raise ValueError("文件不能超过 20 MB")
    suffix = path.suffix.lower()
    if suffix not in ALLOWED_HOSPITAL_SUFFIXES:
        raise ValueError("仅支持 .xls、.xlsx、.xlsm 或 .csv 文件")
    if suffix == ".csv":
        # CSV is decoded and structurally validated by read_import_rows below.
        return
    if not is_zipfile(path):
        # Legacy .xls is parsed through xlrd and converted to a value-only
        # workbook in memory. This does not execute or retain VBA macros.
        if suffix == ".xls":
            workbook = load_excel_path(path, read_only=True, data_only=True)
            workbook.close()
            return
        raise ValueError("文件内容与 Excel 格式不匹配或文件已经损坏")

    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise ValueError("工作簿内部文件数量异常")

            total_uncompressed = 0
            for entry in entries:
                normalized = entry.filename.replace("\\", "/").lower()
                member_path = PurePosixPath(normalized)
                if member_path.is_absolute() or ".." in member_path.parts:
                    raise ValueError("工作簿包含不安全的内部路径")
                if any(part in normalized for part in FORBIDDEN_PACKAGE_PARTS):
                    raise ValueError("工作簿包含宏、外部链接或嵌入对象，不能上传")
                total_uncompressed += entry.file_size
                if total_uncompressed > MAX_UNCOMPRESSED_BYTES:
                    raise ValueError("工作簿解压后体积过大")
                if entry.compress_size > 0 and entry.file_size / entry.compress_size > MAX_COMPRESSION_RATIO:
                    raise ValueError("工作簿压缩比例异常")
    except BadZipFile as exc:
        raise ValueError("无法读取工作簿，请确认文件未损坏") from exc


def inspect_xlsx_package(path: Path) -> None:
    """Backward-compatible name retained for existing callers and tests."""
    inspect_hospital_upload_file(path)


def _append_error(
    errors: list[dict[str, Any]],
    *,
    row_number: int | None,
    sample_id: str | None,
    sample_code: str | None,
    reason: str,
    error_type: str,
) -> None:
    if len(errors) >= MAX_STORED_ERRORS:
        return
    errors.append(
        {
            "row_number": row_number,
            "sample_id": sample_id,
            "sample_code": sample_code,
            "reason": reason,
            "error_type": error_type,
        }
    )


def resolve_workbook_sheet_name(path: Path, requested_sheet_name: str | None) -> str:
    if path.suffix.lower() == ".csv":
        return CSV_SHEET_NAME
    workbook = load_excel_path(path, read_only=True, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
    finally:
        workbook.close()
    if not sheet_names:
        raise ValueError("工作簿中没有工作表")
    requested = (requested_sheet_name or "").strip()
    if requested and requested in sheet_names:
        return requested
    if len(sheet_names) == 1:
        return sheet_names[0]
    if requested:
        raise ValueError(
            f"工作表不存在：{requested}；可用工作表：{'、'.join(sheet_names)}"
        )
    raise ValueError(f"文件包含多个工作表，请填写其中一个：{'、'.join(sheet_names)}")


def _first_row_value(raw_data: dict[str, Any], headers: tuple[str, ...]) -> Any:
    for header in headers:
        value = raw_data.get(header)
        if value is not None and str(value).strip():
            return value
    return None


def _build_code_input(
    headers: list[str],
    row: tuple[Any, ...],
    row_number: int,
    center_code: str,
    allowed_specimen_types: set[str],
) -> SampleCodeInput:
    payload = build_payload(
        headers,
        row,
        row_number,
        allow_nonstandard=True,
        require_sample_code=False,
    )
    raw_data = payload["raw_data"]
    existing_sample_code = payload.get("sample_code")
    if existing_sample_code:
        parsed_code = parse_sample_code(existing_sample_code)
        if parsed_code is not None and parsed_code[0] != center_code:
            raise ValueError(
                f"第 {row_number} 行：文件中已有的平台样本编码属于中心 {parsed_code[0]}，"
                f"与当前账号绑定中心 {center_code} 不一致；医院原始文件无需提供 sample_code 列"
            )
    collection_value = _first_row_value(raw_data, COLLECTION_TIME_HEADERS)
    if collection_value is None:
        raise ValueError(f"第 {row_number} 行：采集时间为空")
    try:
        collection_date = parse_collection_date(collection_value)
    except ValueError as exc:
        raise ValueError(f"第 {row_number} 行：{exc}") from exc
    specimen_type = str(
        _first_row_value(raw_data, SPECIMEN_TYPE_HEADERS) or ""
    ).strip()
    if not specimen_type:
        raise ValueError(f"第 {row_number} 行：样本类型为空")
    if specimen_type not in allowed_specimen_types:
        raise ValueError(
            f"第 {row_number} 行：不支持的样本类型：{specimen_type}"
        )
    return SampleCodeInput(
        row_number=row_number,
        sample_id=payload["sample_id"],
        collection_date=collection_date,
        specimen_type=specimen_type,
    )


def validate_hospital_upload(
    db: Session,
    path: Path,
    sheet_name: str,
    center_code: str,
) -> HospitalValidationResult:
    inspect_hospital_upload_file(path)

    if path.suffix.lower() != ".csv":
        workbook = load_excel_path(path, data_only=False)
        try:
            if len(workbook.sheetnames) > MAX_WORKSHEETS:
                raise ValueError(f"工作簿最多允许 {MAX_WORKSHEETS} 个工作表")
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"工作表不存在：{sheet_name}")
            worksheet = workbook[sheet_name]
            if worksheet.max_row > MAX_ROWS + 1:
                raise ValueError(f"单个工作表最多允许 {MAX_ROWS} 行数据")
            if worksheet.max_column > MAX_COLUMNS:
                raise ValueError(f"单个工作表最多允许 {MAX_COLUMNS} 列")
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        raise ValueError("工作表包含公式，请先转为数值后再上传")
        finally:
            workbook.close()

    center_exists = db.execute(
        text(
            """
            SELECT 1
            FROM centers
            WHERE center_code = :center_code
              AND is_active = true
            """
        ),
        {"center_code": center_code},
    ).scalar_one_or_none()
    if center_exists is None:
        raise ValueError("账号绑定中心不存在或已停用")

    allowed_specimen_types = get_active_specimen_type_names(db)

    headers, rows, _ = read_import_rows(path, sheet_name)
    if len(headers) > MAX_COLUMNS:
        raise ValueError(f"单个文件最多允许 {MAX_COLUMNS} 列")
    if path.suffix.lower() == ".csv":
        for row_number, row in enumerate(rows, start=2):
            for value in row:
                if not isinstance(value, str):
                    continue
                candidate = value.lstrip()
                starts_formula = candidate.startswith(("=", "@", "\t", "\r"))
                starts_named_formula = (
                    len(candidate) > 1
                    and candidate[0] in {"+", "-"}
                    and candidate[1].isalpha()
                )
                if starts_formula or starts_named_formula:
                    raise ValueError(
                        f"CSV 第 {row_number} 行包含可能被电子表格软件执行的公式内容，请先转为普通文本"
                    )
    nonempty_headers = [header for header in headers if header]
    if len(nonempty_headers) != len(set(nonempty_headers)):
        raise ValueError("表头存在重复字段，请使用标准模板重新整理")
    required_header_groups = (
        ("医院原始序号", SAMPLE_ID_HEADERS),
        ("采集时间", COLLECTION_TIME_HEADERS),
        ("样本类型", SPECIMEN_TYPE_HEADERS),
    )
    missing_groups = [
        label
        for label, aliases in required_header_groups
        if not any(header in headers for header in aliases)
    ]
    if missing_groups:
        raise ValueError("缺少必填表头：" + "、".join(missing_groups))

    nonempty_rows = [
        (row_number, row)
        for row_number, row in enumerate(rows, start=2)
        if any(value is not None and str(value).strip() for value in row)
    ]
    if not nonempty_rows:
        raise ValueError("工作表中没有可提交的数据")
    if len(nonempty_rows) > MAX_ROWS:
        raise ValueError(f"单次最多允许 {MAX_ROWS} 行数据")

    errors: list[dict[str, Any]] = []
    all_error_rows: set[int] = set()
    parsed_rows: list[SampleCodeInput] = []
    id_rows: dict[str, list[int]] = {}

    for row_number, row in nonempty_rows:
        sample_id = None
        try:
            code_input = _build_code_input(
                headers,
                row,
                row_number,
                center_code,
                allowed_specimen_types,
            )
            sample_id = code_input.sample_id
            parsed_rows.append(code_input)
            id_rows.setdefault(sample_id, []).append(row_number)
        except ValueError as exc:
            all_error_rows.add(row_number)
            _append_error(
                errors,
                row_number=row_number,
                sample_id=sample_id,
                sample_code=None,
                reason=str(exc),
                error_type="validation",
            )

    duplicate_in_file = {
        sample_id for sample_id, row_numbers in id_rows.items() if len(row_numbers) > 1
    }
    for sample_id in sorted(duplicate_in_file):
        for row_number in id_rows[sample_id]:
            all_error_rows.add(row_number)
            _append_error(
                errors,
                row_number=row_number,
                sample_id=sample_id,
                sample_code=None,
                reason=f"第 {row_number} 行：医院原始序号在文件内重复：{sample_id}",
                error_type="duplicate_in_file",
            )

    total_rows = len(nonempty_rows)
    error_rows = len(all_error_rows)
    valid_rows = total_rows - error_rows
    if len(errors) >= MAX_STORED_ERRORS and error_rows > MAX_STORED_ERRORS:
        errors[-1] = {
            "row_number": None,
            "sample_id": None,
            "sample_code": None,
            "reason": f"错误数量较多，仅保留前 {MAX_STORED_ERRORS} 条，请下载报告后分批修正",
            "error_type": "truncated",
        }

    summary = {
        "template_version": TEMPLATE_VERSION,
        "center_code": center_code,
        "sheet_name": sheet_name,
        "required_fields": ["医院原始序号", "采集时间", "样本类型"],
        "file_duplicate_original_ids": len(duplicate_in_file),
        "date_counts": {
            day: sum(1 for row in parsed_rows if row.collection_date.isoformat() == day)
            for day in sorted({row.collection_date.isoformat() for row in parsed_rows})
        },
        "specimen_type_counts": {
            specimen_type: sum(1 for row in parsed_rows if row.specimen_type == specimen_type)
            for specimen_type in sorted({row.specimen_type for row in parsed_rows})
        },
        "error_report_truncated": error_rows > len(errors),
    }
    return HospitalValidationResult(
        total_rows=total_rows,
        valid_rows=valid_rows,
        error_rows=error_rows,
        new_rows=valid_rows,
        duplicate_rows=len(duplicate_in_file),
        errors=errors,
        summary=summary,
    )


def build_hospital_code_assignments(
    db: Session,
    path: Path,
    sheet_name: str,
    center_code: str,
    *,
    lock_center: bool = False,
) -> tuple[list[SampleCodeAssignment], dict[str, int]]:
    headers, rows, _ = read_import_rows(path, sheet_name)
    allowed_specimen_types = get_active_specimen_type_names(db)
    code_inputs = [
        _build_code_input(
            headers,
            row,
            row_number,
            center_code,
            allowed_specimen_types,
        )
        for row_number, row in enumerate(rows, start=2)
        if any(value is not None and str(value).strip() for value in row)
    ]
    return allocate_sample_codes(
        db,
        center_code,
        code_inputs,
        lock_center=lock_center,
        allowed_specimen_types=allowed_specimen_types,
    )


def serialize_review_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value)


def build_review_headers(headers: list[str]) -> tuple[list[str], int]:
    preserved_headers = [
        "原表_sample_code" if header.strip().lower() == REVIEW_SAMPLE_CODE_HEADER else header
        for header in headers
    ]
    code_insert_index = preserved_headers.index("序") + 1 if "序" in preserved_headers else 0
    preview_headers = list(preserved_headers)
    preview_headers.insert(code_insert_index, REVIEW_SAMPLE_CODE_HEADER)
    preview_headers.extend(REVIEW_PREVIEW_RESULT_HEADERS)
    return preview_headers, code_insert_index


def build_hospital_submission_preview(
    path: Path,
    sheet_name: str,
    *,
    limit: int | None = 200,
) -> tuple[list[str], list[dict[str, Any]], list[int], int]:
    headers, rows, _ = read_import_rows(path, sheet_name)
    sensitive_columns = [
        index for index, header in enumerate(headers) if header in SENSITIVE_REVIEW_HEADERS
    ]
    preview_rows: list[dict[str, Any]] = []
    total_rows = 0
    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        total_rows += 1
        if limit is not None and len(preview_rows) >= limit:
            continue
        preview_rows.append(
            {
                "row_number": row_number,
                "values": [
                    serialize_review_value(row[index] if index < len(row) else None)
                    for index in range(len(headers))
                ],
            }
        )
    return list(headers), preview_rows, sensitive_columns, total_rows


def build_hospital_review_preview(
    path: Path,
    sheet_name: str,
    assignments: list[SampleCodeAssignment],
    *,
    limit: int | None = 200,
) -> tuple[list[str], list[dict[str, Any]], list[int], int]:
    headers, rows, _ = read_import_rows(path, sheet_name)
    assignment_by_row = {assignment.row_number: assignment for assignment in assignments}
    preview_headers, code_insert_index = build_review_headers(headers)
    sensitive_columns = [
        index + (1 if index >= code_insert_index else 0)
        for index, header in enumerate(headers)
        if header in SENSITIVE_REVIEW_HEADERS
    ]
    preview_rows: list[dict[str, Any]] = []
    total_rows = 0
    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        total_rows += 1
        if limit is not None and len(preview_rows) >= limit:
            continue
        assignment = assignment_by_row.get(row_number)
        original_values = [
            serialize_review_value(row[index] if index < len(row) else None)
            for index in range(len(headers))
        ]
        original_values.insert(
            code_insert_index,
            assignment.sample_code if assignment else None,
        )
        preview_rows.append(
            {
                "row_number": row_number,
                "values": [
                    *original_values,
                    "通过" if assignment else "异常",
                    "" if assignment else "未生成候选编码",
                ],
            }
        )
    return preview_headers, preview_rows, sensitive_columns, total_rows


def build_hospital_review_preview_workbook(
    path: Path,
    sheet_name: str,
    assignments: list[SampleCodeAssignment],
) -> bytes:
    headers, rows, _ = read_import_rows(path, sheet_name)
    assignment_by_row = {assignment.row_number: assignment for assignment in assignments}
    preview_headers, code_insert_index = build_review_headers(headers)
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("审核预览")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers) + 3)}1"

    header_fill = PatternFill("solid", fgColor="1D4ED8")
    header_cells = []
    for header in preview_headers:
        cell = WriteOnlyCell(worksheet, value=header)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cells.append(cell)
    worksheet.append(header_cells)

    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        assignment = assignment_by_row.get(row_number)
        original_values = [
            serialize_review_value(row[index] if index < len(row) else None)
            for index in range(len(headers))
        ]
        original_values.insert(
            code_insert_index,
            assignment.sample_code if assignment else None,
        )
        worksheet.append(
            [
                *original_values,
                "通过" if assignment else "异常",
                "" if assignment else "未生成候选编码",
            ]
        )

    notes = workbook.create_sheet("审核说明")
    notes.append(["项目", "说明"])
    notes.append(["原始工作表", sheet_name])
    notes.append(["sample_code", "系统生成的候选正式编码；使用本批次正式导入时会锁定中心并重新校验"])
    notes.append(["校验结果", "下载时已重新校验原文件；正式导入前仍会再次校验"])
    notes.append(["敏感数据", "文件可能包含个人敏感信息，仅限获授权人员使用并按规定留存"])

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_hospital_upload_template() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "样本数据"

    core_headers = list(CORE_FIELD_MAP.keys())
    headers = [
        "序",
        "采集时间",
        "标本",
        "条码号",
        "实验号",
        "病人号",
        "姓名",
        "身份证号",
        "电话",
        "地址",
        *[
            header
            for header in core_headers
            if header not in {"序", "采集时间", "标本", "条码号", "实验号", "病人号"}
        ],
        "WBC",
    ]
    worksheet.append(headers)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{worksheet.cell(1, len(headers)).coordinate}"
    header_fill = PatternFill("solid", fgColor="2563EB")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, header in enumerate(headers, start=1):
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = min(
            max(len(header) * 2 + 4, 12),
            24,
        )

    instructions = workbook.create_sheet("填写说明")
    instruction_rows = [
        ("模板版本", TEMPLATE_VERSION),
        ("必填字段", "医院原始序号（如“序”）、采集时间、样本类型（如“标本”）"),
        ("样本编码", "医院无需填写 sample_code；内部审核时按账号中心和采集日期统一生成"),
        ("可选字段", "除必填字段外可保留医院原有字段；系统无法识别的字段仍保存到原始 JSON"),
        ("重复规则", "文件内医院原始序号不能重复；正式编码在审核入库前再次查重"),
        ("文件要求", "支持 .xls / .xlsx / .xlsm / .csv；不得包含宏、公式、外部链接或嵌入对象"),
        ("检验字段", "可从 WBC 列开始继续增加双方已确认的检验项目列"),
        ("处理流程", "上传校验通过后提交审核；审核通过可仅留档，也可由内部管理员按需选择导入系统"),
    ]
    for row in instruction_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 18
    instructions.column_dimensions["B"].width = 86
    for cell in instructions[1]:
        cell.font = Font(bold=True)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
