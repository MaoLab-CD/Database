from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.services.sample_codes import DerivedSampleRule, parse_sample_code_parts
from app.services.specimen_types import get_confirmed_specimen_code_rules
from app.services.workbook_loader import load_excel_workbook


CODE_HEADER = "sample_code"
CODE_HEADERS = (CODE_HEADER, "样本编码", "正式编码", "贴标编码")
SEQUENCE_HEADER = "序"
COLLECTION_TIME_HEADERS = ("采集时间", "采样时间", "采集日期", "采样日期")
MAX_SERIAL = 9999
WENJIANG_CENTER_CODE = "510115-001"
WENJIANG_SEQUENCE_DATE_PATTERN = re.compile(
    r"^(?P<month>\d{2})(?P<day>\d{2})(?P<serial>\d{2})$"
)


@dataclass(frozen=True)
class SampleCodeInput:
    row_number: int
    sample_id: str
    collection_date: date
    specimen_type: str
    existing_sample_code: str | None = None


@dataclass(frozen=True)
class SampleCodeAssignment:
    row_number: int
    sample_id: str
    collection_date: str
    specimen_type: str
    sample_code: str
    assignment_source: str = "generated"


def _identity_key(row: SampleCodeInput) -> tuple[str, str, str]:
    return (
        row.sample_id,
        row.specimen_type,
        row.collection_date.strftime("%y%m%d"),
    )


def get_existing_sample_codes_by_identity(
    db: Session,
    center_code: str,
    rows: list[SampleCodeInput],
) -> dict[tuple[str, str, str], list[str]]:
    if not rows:
        return {}
    result_rows = (
        db.execute(
            text(
                """
                SELECT sample_id,
                       specimen_type,
                       substring(sample_seq FROM 1 FOR 6) AS date_prefix,
                       sample_code
                FROM samples
                WHERE is_deleted = false
                  AND center_code = :center_code
                  AND sample_id = ANY(:sample_ids)
                  AND specimen_type = ANY(:specimen_types)
                  AND substring(sample_seq FROM 1 FOR 6) = ANY(:date_prefixes)
                  AND sample_seq ~ '^[0-9]{10}[A-Za-z]?$'
                ORDER BY id
                """
            ),
            {
                "center_code": center_code,
                "sample_ids": sorted({row.sample_id for row in rows}),
                "specimen_types": sorted({row.specimen_type for row in rows}),
                "date_prefixes": sorted(
                    {row.collection_date.strftime("%y%m%d") for row in rows}
                ),
            },
        )
        .mappings()
        .all()
    )
    matches: dict[tuple[str, str, str], list[str]] = defaultdict(list)
    for result_row in result_rows:
        sample_code = str(result_row["sample_code"] or "").strip()
        if sample_code:
            matches[
                (
                    str(result_row["sample_id"]),
                    str(result_row["specimen_type"]),
                    str(result_row["date_prefix"]),
                )
            ].append(sample_code)
    return dict(matches)


def get_sample_code_owners(
    db: Session,
    sample_codes: set[str],
) -> dict[str, tuple[str, str, str]]:
    if not sample_codes:
        return {}
    rows = (
        db.execute(
            text(
                """
                SELECT sample_code,
                       sample_id,
                       specimen_type,
                       substring(sample_seq FROM 1 FOR 6) AS date_prefix
                FROM samples
                WHERE is_deleted = false
                  AND sample_code = ANY(:sample_codes)
                """
            ),
            {"sample_codes": sorted(sample_codes)},
        )
        .mappings()
        .all()
    )
    return {
        str(row["sample_code"]): (
            str(row["sample_id"]),
            str(row["specimen_type"]),
            str(row["date_prefix"]),
        )
        for row in rows
    }


def allocate_sample_codes(
    db: Session,
    center_code: str,
    rows: list[SampleCodeInput],
    *,
    lock_center: bool = False,
    allowed_specimen_types: set[str] | None = None,
    code_rules: tuple[DerivedSampleRule, ...] | None = None,
    reuse_existing: bool = False,
) -> tuple[list[SampleCodeAssignment], dict[str, int]]:
    center_sql = """
        SELECT id
        FROM centers
        WHERE center_code = :center_code
          AND is_active = true
    """
    if lock_center:
        center_sql += " FOR UPDATE"
    center_id = db.execute(text(center_sql), {"center_code": center_code}).scalar_one_or_none()
    if center_id is None:
        raise ValueError("中心不存在或已停用")
    if not rows:
        raise ValueError("没有可生成编码的数据行")

    if code_rules is None:
        code_rules = get_confirmed_specimen_code_rules(db)
    rule_by_type = {rule.specimen_type: rule for rule in code_rules}
    if allowed_specimen_types is None:
        allowed_specimen_types = set(rule_by_type)

    seen_identities: dict[tuple[str, str, str], int] = {}
    seen_provided_codes: dict[str, int] = {}
    for row in rows:
        identity = _identity_key(row)
        previous_row = seen_identities.get(identity)
        if previous_row is not None:
            raise ValueError(
                f"第 {row.row_number} 行与第 {previous_row} 行是同一中心、类型、采集日期和序，"
                "请先删除重复行"
            )
        seen_identities[identity] = row.row_number
        if row.existing_sample_code:
            previous_code_row = seen_provided_codes.get(row.existing_sample_code)
            if previous_code_row is not None:
                raise ValueError(
                    f"第 {row.row_number} 行与第 {previous_code_row} 行使用了相同样本编码："
                    f"{row.existing_sample_code}"
                )
            seen_provided_codes[row.existing_sample_code] = row.row_number

    for row in rows:
        if row.specimen_type not in allowed_specimen_types:
            raise ValueError(
                f"第 {row.row_number} 行：不支持的样本类型：{row.specimen_type}"
            )
        if row.specimen_type not in rule_by_type:
            raise ValueError(
                f"第 {row.row_number} 行：样本类型“{row.specimen_type}”的编码规则尚未确认"
            )

    existing_by_identity = (
        get_existing_sample_codes_by_identity(db, center_code, rows)
        if reuse_existing
        else {}
    )
    owners_by_code = (
        get_sample_code_owners(
            db,
            {row.existing_sample_code for row in rows if row.existing_sample_code},
        )
        if reuse_existing
        else {}
    )

    resolved: list[tuple[SampleCodeInput, str | None, str]] = []
    date_prefixes = {row.collection_date.strftime("%y%m%d") for row in rows}
    counters: dict[str, int] = defaultdict(
        int,
        get_database_max_serials(db, center_code, date_prefixes),
    )

    for row in rows:
        identity = _identity_key(row)
        matches = list(dict.fromkeys(existing_by_identity.get(identity, [])))
        if len(matches) > 1:
            raise ValueError(
                f"第 {row.row_number} 行：数据库中找到多个同中心、类型、采集日期和序的样本，"
                "无法安全复用编码"
            )

        if row.existing_sample_code:
            parts = parse_sample_code_parts(row.existing_sample_code, code_rules)
            if parts is None:
                raise ValueError(
                    f"第 {row.row_number} 行：已有样本编码格式不正确或后缀未确认："
                    f"{row.existing_sample_code}"
                )
            if parts.center_code != center_code:
                raise ValueError(
                    f"第 {row.row_number} 行：已有样本编码属于中心 {parts.center_code}，"
                    f"与当前中心 {center_code} 不一致"
                )
            if parts.specimen_type != row.specimen_type:
                raise ValueError(
                    f"第 {row.row_number} 行：已有样本编码对应类型为 {parts.specimen_type}，"
                    f"与当前选择的 {row.specimen_type} 不一致"
                )
            date_prefix = row.collection_date.strftime("%y%m%d")
            if not parts.base_seq.startswith(date_prefix):
                raise ValueError(
                    f"第 {row.row_number} 行：已有样本编码日期与采集日期不一致："
                    f"{row.existing_sample_code}"
                )
            owner = owners_by_code.get(row.existing_sample_code)
            if owner is not None and owner != identity:
                raise ValueError(
                    f"第 {row.row_number} 行：已有样本编码已属于数据库中的其他样本："
                    f"{row.existing_sample_code}"
                )
            if matches and matches[0] != row.existing_sample_code:
                raise ValueError(
                    f"第 {row.row_number} 行：该样本数据库编码为 {matches[0]}，"
                    f"但 Excel 中是 {row.existing_sample_code}"
                )
            counters[date_prefix] = max(
                counters[date_prefix],
                int(parts.base_seq[6:10]),
            )
            resolved.append((row, row.existing_sample_code, "provided"))
            continue

        if matches:
            parts = parse_sample_code_parts(matches[0], code_rules)
            if parts is None or parts.specimen_type != row.specimen_type:
                raise ValueError(
                    f"第 {row.row_number} 行：数据库已有编码不符合当前已确认规则：{matches[0]}"
                )
            date_prefix = row.collection_date.strftime("%y%m%d")
            counters[date_prefix] = max(
                counters[date_prefix],
                int(parts.base_seq[6:10]),
            )
            resolved.append((row, matches[0], "reused"))
            continue

        resolved.append((row, None, "generated"))

    counts: dict[str, int] = defaultdict(int)
    assignments: list[SampleCodeAssignment] = []
    for row, resolved_code, assignment_source in resolved:
        date_prefix = row.collection_date.strftime("%y%m%d")
        if resolved_code is None:
            counters[date_prefix] += 1
            if counters[date_prefix] > MAX_SERIAL:
                raise ValueError(f"日期 {date_prefix} 的流水号超过 {MAX_SERIAL}")
            suffix = rule_by_type[row.specimen_type].suffix
            sample_seq = f"{date_prefix}{counters[date_prefix]:04d}{suffix}"
            resolved_code = f"{center_code}-{sample_seq}"
        assignments.append(
            SampleCodeAssignment(
                row_number=row.row_number,
                sample_id=row.sample_id,
                collection_date=row.collection_date.isoformat(),
                specimen_type=row.specimen_type,
                sample_code=resolved_code,
                assignment_source=assignment_source,
            )
        )
        counts[row.collection_date.isoformat()] += 1
    return assignments, dict(sorted(counts.items()))


def parse_collection_date(value: Any) -> date:
    if value is None or value == "":
        raise ValueError("采集时间为空")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    raw = str(value).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y/%m/%d",
        "%Y.%m.%d %H:%M:%S",
        "%Y.%m.%d %H:%M",
        "%Y.%m.%d",
    ):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"无法解析采集时间：{raw}")


def normalize_header(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def find_collection_time_column(worksheet: Any) -> int:
    headers = {
        normalize_header(cell.value): cell.column
        for cell in worksheet[1]
        if normalize_header(cell.value)
    }
    for header in COLLECTION_TIME_HEADERS:
        if header in headers:
            return headers[header]
    raise ValueError(
        "未找到采集时间列，支持的表头为：" + "、".join(COLLECTION_TIME_HEADERS)
    )


def get_database_max_serials(
    db: Session,
    center_code: str,
    date_prefixes: set[str],
) -> dict[str, int]:
    if not date_prefixes:
        return {}
    rows = (
        db.execute(
            text(
                """
                SELECT substring(sample_seq FROM 1 FOR 6) AS date_prefix,
                       max(substring(sample_seq FROM 7 FOR 4)::int) AS max_serial
                FROM samples
                WHERE is_deleted = false
                  AND center_code = :center_code
                  AND substring(sample_seq FROM 1 FOR 6) = ANY(:date_prefixes)
                  AND sample_seq ~ '^[0-9]{10}[A-Za-z]?$'
                GROUP BY substring(sample_seq FROM 1 FOR 6)
                """
            ),
            {"center_code": center_code, "date_prefixes": list(date_prefixes)},
        )
        .mappings()
        .all()
    )
    return {str(row["date_prefix"]): int(row["max_serial"] or 0) for row in rows}


def build_excel_code_inputs(
    worksheet: Any,
    *,
    sequence_column: int,
    collection_column: int,
    code_column: int,
    center_code: str,
    sample_type: str,
) -> list[SampleCodeInput]:
    parsed_rows: list[tuple[int, str, date | None]] = []
    errors: list[str] = []

    for row_index in range(2, worksheet.max_row + 1):
        values = [
            worksheet.cell(row=row_index, column=column).value
            for column in range(1, worksheet.max_column + 1)
            if column != code_column
        ]
        if not any(value is not None and str(value).strip() for value in values):
            continue

        sample_id = normalize_header(
            worksheet.cell(row=row_index, column=sequence_column).value
        )
        if not sample_id:
            errors.append(f"第 {row_index} 行：序为空")
            continue

        collection_value = worksheet.cell(
            row=row_index,
            column=collection_column,
        ).value
        if collection_value is None or not str(collection_value).strip():
            parsed_rows.append((row_index, sample_id, None))
            continue
        try:
            collection_date = parse_collection_date(collection_value)
        except ValueError as exc:
            errors.append(f"第 {row_index} 行：{exc}")
            continue
        parsed_rows.append((row_index, sample_id, collection_date))

    if errors:
        raise ValueError(_format_excel_code_errors(errors))
    if not parsed_rows:
        raise ValueError("工作表中没有可生成编码的数据行")

    missing_rows = [row for row in parsed_rows if row[2] is None]
    if not missing_rows:
        return [
            SampleCodeInput(
                row_number=row_index,
                sample_id=sample_id,
                collection_date=collection_date,
                specimen_type=sample_type,
                existing_sample_code=(
                    normalize_header(worksheet.cell(row=row_index, column=code_column).value)
                    or None
                ),
            )
            for row_index, sample_id, collection_date in parsed_rows
            if collection_date is not None
        ]

    if center_code != WENJIANG_CENTER_CODE:
        raise ValueError(
            _format_excel_code_errors(
                [f"第 {row_index} 行：采集时间为空" for row_index, _, _ in missing_rows]
            )
        )

    known_years = {
        collection_date.year
        for _, _, collection_date in parsed_rows
        if collection_date is not None
    }
    if not known_years:
        raise ValueError(
            "采集时间校验失败：温江区人民医院仅能在表内至少有一条有效采集时间时，"
            "根据“序”的 MMDD 推断缺失行日期"
        )
    if len(known_years) != 1:
        raise ValueError(
            "采集时间校验失败：表内采集时间跨越多个年份，无法安全地根据“序”推断缺失行年份"
        )
    inferred_year = next(iter(known_years))

    code_inputs: list[SampleCodeInput] = []
    fallback_errors: list[str] = []
    for row_index, sample_id, collection_date in parsed_rows:
        match = WENJIANG_SEQUENCE_DATE_PATTERN.fullmatch(sample_id)
        if match is None:
            fallback_errors.append(
                f"第 {row_index} 行：温江兼容模式要求“序”为 6 位 MMDDNN：{sample_id}"
            )
            continue
        serial = int(match.group("serial"))
        if serial < 1:
            fallback_errors.append(
                f"第 {row_index} 行：“序”的当日流水不能为 00：{sample_id}"
            )
            continue
        try:
            inferred_date = date(
                inferred_year,
                int(match.group("month")),
                int(match.group("day")),
            )
        except ValueError:
            fallback_errors.append(
                f"第 {row_index} 行：“序”包含无效日期：{sample_id}"
            )
            continue
        if collection_date is not None and collection_date != inferred_date:
            fallback_errors.append(
                f"第 {row_index} 行：采集日期 {collection_date.isoformat()} 与“序” {sample_id} 的日期不一致"
            )
            continue
        code_inputs.append(
            SampleCodeInput(
                row_number=row_index,
                sample_id=sample_id,
                collection_date=collection_date or inferred_date,
                specimen_type=sample_type,
                existing_sample_code=(
                    normalize_header(worksheet.cell(row=row_index, column=code_column).value)
                    or None
                ),
            )
        )

    if fallback_errors:
        raise ValueError(_format_excel_code_errors(fallback_errors))
    return code_inputs


def _format_excel_code_errors(errors: list[str]) -> str:
    preview = "；".join(errors[:20])
    suffix = f"；另有 {len(errors) - 20} 行" if len(errors) > 20 else ""
    return f"采集时间校验失败：{preview}{suffix}"


def add_sample_codes_to_excel(
    content: bytes,
    sheet_name: str,
    center_code: str,
    sample_type: str,
    db: Session,
    keep_vba: bool = False,
    source_suffix: str = ".xlsx",
) -> tuple[bytes, dict[str, int], dict[str, int], list[dict[str, Any]]]:
    workbook = load_excel_workbook(
        content,
        source_suffix,
        keep_vba=keep_vba,
    )
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"未找到工作表：{sheet_name}")
        worksheet = workbook[sheet_name]

        code_columns = [
            cell.column
            for cell in worksheet[1]
            if normalize_header(cell.value) in CODE_HEADERS
        ]
        if len(code_columns) > 1:
            raise ValueError("检测到多个样本编码列，请只保留一个后再生成")
        if code_columns:
            code_column = code_columns[0]
        else:
            worksheet.insert_cols(2)
            worksheet.cell(row=1, column=2).value = CODE_HEADER
            code_column = 2

        sequence_column = next(
            (
                cell.column
                for cell in worksheet[1]
                if normalize_header(cell.value) == SEQUENCE_HEADER
            ),
            None,
        )
        if sequence_column is None:
            raise ValueError("未找到序列，请先在 Excel 中补充表头为“序”的列")

        collection_column = find_collection_time_column(worksheet)
        code_inputs = build_excel_code_inputs(
            worksheet,
            sequence_column=sequence_column,
            collection_column=collection_column,
            code_column=code_column,
            center_code=center_code,
            sample_type=sample_type,
        )
        code_rules = get_confirmed_specimen_code_rules(
            db,
            batch_code_only=True,
        )
        assignments, counts = allocate_sample_codes(
            db,
            center_code,
            code_inputs,
            allowed_specimen_types={rule.specimen_type for rule in code_rules},
            code_rules=code_rules,
            reuse_existing=True,
        )
        source_counts: dict[str, int] = defaultdict(int)
        preview_rows: list[dict[str, Any]] = []

        for assignment in assignments:
            source_counts[assignment.assignment_source] += 1
            cell = worksheet.cell(row=assignment.row_number, column=code_column)
            cell.value = assignment.sample_code
            cell.alignment = Alignment(horizontal="left")
            if len(preview_rows) < 100:
                preview_rows.append(
                    {
                        "row_number": assignment.row_number,
                        "collection_date": assignment.collection_date,
                        "sample_code": cell.value,
                        "assignment_source": assignment.assignment_source,
                    }
                )

        header_cell = worksheet.cell(row=1, column=code_column)
        header_cell.font = Font(bold=True)
        header_cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        header_cell.alignment = Alignment(horizontal="center")
        worksheet.column_dimensions[get_column_letter(code_column)].width = 28

        output = BytesIO()
        workbook.save(output)
        return output.getvalue(), counts, dict(source_counts), preview_rows
    finally:
        workbook.close()
